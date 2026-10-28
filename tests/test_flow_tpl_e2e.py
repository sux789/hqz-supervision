#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""v2 流程模板（设计值 / 流程路由 / 参数 / 流程-*）端到端测试（v0.38 doc/030、v0.39 doc/031）

用法：python tests/test_flow_tpl_e2e.py      # 全绿退出码 0

模板结构变了（一列一因子 → 一行一环节），App 的行为也跟着变：
  T1  上传 → flows.schema=flow_tpl/1；5 条路由；设计值表头映射成 std
  T1c 「参数」页驱动搜索栏（项目类型→标段→乡→村→小班，级联顺序）
  T1d 「参数」页带出照片配置（目录/水印/文件名/压缩）
  T2  不带 flow → 默认给「当前该做的那条」，其余按路由顺序放进 flow_list
  T2g 独立流程（日常巡护）排最后、无前置、可多次，且不抢默认流程
  T3  前置未完成 → flow_locked（GET 可预览，POST 409）
  T4  不通过 / 选了带 →备注 的选项 → 备注不足 2 个汉字 400
  T5  必填缺项 400；数值 -5 → 400（服务端权威，C13）
  T6  整条流程一次提交 → 全部环节同一 rec_no
  T7  通过后下一条解锁（整地→回塘底肥→栽植→完工）
  T8  自动(公式)由服务端算：90/(90+10)×100 = 90
  T9  独立流程先交不阻塞主线；可重复流程每次 rec_no 递增；带 rec_no = 改那一次
  T10 v2 工作簿走老 /flow 提交 → 400（不许绕过整条校验）
  T11 模板附件列：上传 ≥3 份（图片≥1 / PDF≥1），坏标签/坏类型/缺类型都 400
  T12 附件落库与回显：下载鉴权、rec_no 绑定、跨记录复用拒绝

模板放在仓库 data/ 下（可用 TPL_V2 覆盖）；找不到整段 SKIP，不算失败。
"""
import io
import os
import shutil
import sqlite3
import sys
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

import main as _M  # noqa: E402
from main import create_app  # noqa: E402

ADMIN = ('雷华雄', 'lhx123')

CANDIDATES = [
    os.environ.get('TPL_V2'),
    str(BASE / 'data' / 'tpl-人工造林-验收流程模板.xlsx'),
]
TPL = next((Path(p) for p in CANDIDATES if p and Path(p).exists()), None)
if TPL is None:
    TPL = next(iter(sorted((BASE / 'data').rglob('*人工造林*验收流程模板*.xlsx'))), None)

ROUTES = ['流程-整地', '流程-回塘底肥', '流程-栽植撑杆覆膜', '流程-完工验收', '流程-日常巡护']

_pass = _fail = 0


def ok(name, cond, detail=''):
    global _pass, _fail
    _pass, _fail = _pass + bool(cond), _fail + (not cond)
    print(f'  {"✅" if cond else "❌"} {name}  {detail}')


def batch(cli, H, wid, ridx, flow, steps, rec_no=None, geo=None):
    body = {'row': ridx, 'flow': flow, 'steps': steps}
    if rec_no is not None:
        body['rec_no'] = rec_no
    if geo is not None:
        body['geo'] = geo
    return cli.post(f'/api/workbooks/{wid}/flow/batch', headers=H, json=body)


def one_step(no, **values):
    return {'no': no, 'values': values}


PNG = b'\x89PNG\r\n\x1a\n' + b'\x00' * 40
JPG = b'\xff\xd8\xff' + b'\x00' * 40
PDF = b'%PDF-1.4\n' + b'\x00' * 40


def try_att(cli, H, wid, ridx, flow, step, label, name, data):
    """上传一个附件（原始响应，便于断言失败分支）。"""
    return cli.post(
        f'/api/workbooks/{wid}/flow/att', headers=H,
        data={'row': str(ridx), 'flow': flow, 'step': str(step), 'label': label,
              'file': (io.BytesIO(data), name)},
        content_type='multipart/form-data')


def att_entry(cli, H, wid, ridx, flow, step, label, files):
    """上传一组文件 → 提交用的 att 条目（ids 由服务端发）。"""
    ids, names, kinds = [], [], []
    for name, data in files:
        r = try_att(cli, H, wid, ridx, flow, step, label, name, data)
        j = r.get_json() or {}
        assert r.status_code == 200, f'附件上传失败 {name}: {r.status_code} {j}'
        ids.append(j['id'])
        names.append(j.get('name') or name)
        kinds.append(j.get('kind') or 'image')
    return {'kind': 'upload', 'label': label, 'min': 3, 'count': len(ids),
            'names': names, 'types': kinds, 'ids': ids}


def full_steps(values_by_step, att=None):
    """{步骤号: {因子: 值}} → 批量提交的 steps 数组。"""
    return [{'no': no, 'values': vals, 'att': (att or {}).get(no, [])}
            for no, vals in values_by_step.items()]


LAND_PASS = {
    1: {'核实施工面积': '正确'},
    2: {'核实密度': '正确'},
    3: {'核实株行距': '3.0×4.0'},
    4: {'核实整地方式': '穴状整地'},
    5: {'核实整地规格': '60×60×50'},
    6: {'核实总株数': 6076},
    8: {'是否通过': '通过'},
}
BASE_PASS = {
    1: {'核实施工面积': '正确'},
    2: {'有机肥1.5kg/株': '正确'},
    3: {'普钙0.2kg/株': '正确'},
    4: {'回塘质量': '正确'},
    6: {'是否通过': '通过'},
}
PLANT_PASS = {
    1: {'核实栽植面积': '正确'},
    2: {'核实造林树种': '正确'},
    3: {'核实混交方式': '正确'},
    4: {'核实混交比例': '正确'},
    5: {'核实栽植总株数': 6076},
    6: {'核实树种1株数': 3646},
    7: {'核实树种2株数': 2430},
    8: {'核实撑杆数量': 6076},
    9: {'撑杆规格': '正确'},
    10: {'核实覆膜数量': 6076},
    11: {'覆膜质量': '正确'},
    13: {'是否通过': '通过'},
}
DONE_PASS = {
    1: {'核实栽植面积': '正确'},
    2: {'核实造林树种': '正确'},
    3: {'核实混交方式': '正确'},
    4: {'核实混交比例': '正确'},
    5: {'核实栽植总株数': 6076},
    6: {'核实树种1株数': 3646},
    7: {'核实树种2株数': 2430},
    8: {'栽植成活率': 92},
    9: {'核实撑杆数量': 6076},
    10: {'撑杆规格': '正确'},
    11: {'核实覆膜数量': 6076},
    12: {'覆膜质量': '正确'},
    13: {'样地调查成活株数': 90},
    14: {'样地调查死亡株数': 10},
    17: {'是否通过': '通过'},
}
PATROL_PASS = {1: {'巡护情况': ['林地清理、整地巡护']}, 3: {'是否通过': '通过'}}


def main():
    if TPL is None:
        print('  ⏭  找不到 v2 模板（tpl-人工造林-验收流程模板.xlsx）→ 整段 SKIP')
        return 0

    app = create_app()
    cli = app.test_client()
    r = cli.post('/api/login', json={'username': ADMIN[0], 'password': ADMIN[1]})
    tok = (r.get_json() or {}).get('token', '')
    if r.status_code != 200 or not tok:
        print(f'  ⏭  管理员登录失败（HTTP {r.status_code}，可能改过密码）→ 整段 SKIP')
        return 0
    H = {'X-Sup-Token': tok}

    buf = io.BytesIO(TPL.read_bytes())
    r = cli.post('/api/workbooks', data={'file': (buf, TPL.name)},
                 headers=H, content_type='multipart/form-data')
    j = r.get_json() or {}
    wid = j.get('id')
    ok('T0 v2 模板能上传', r.status_code == 200 and bool(wid), str(j)[:120])
    if not wid:
        return 1

    try:
        g = cli.get(f'/api/workbooks/{wid}', headers=H).get_json()
        fl = g.get('flows') or {}
        hdrs, rows = g.get('headers') or [], g.get('rows') or []
        ok('T1 落库为 v2 schema（flow_tpl/1）且路由 5 条',
           fl.get('schema') == 'flow_tpl/1' and len(fl.get('routes') or []) == 5,
           f"{fl.get('schema')} {len(fl.get('routes') or [])} 条")
        ok('T1b 设计值表被认作数据表（149 小班）',
           g.get('sheet_name') == '设计值' and len(rows) == 149,
           f"{g.get('sheet_name')} {len(rows)} 行")
        ok('T1c「参数」页驱动搜索栏：项目类型→标段→乡→村→小班（级联顺序）',
           (g.get('config') or {}).get('搜索选项')
           == '项目类型|select;标段|select;乡（镇）|select;村（社区）|select;小班号|search'
           and g.get('key_column') == '小班号',
           str((g.get('config') or {}).get('搜索选项')))
        cfg = g.get('config') or {}
        ok('T1d 照片参数进模板：本机目录/水印/文件名/压缩',
           cfg.get('目录') == '{{项目类型}}/{{标段}}/{{小班号}}'
           and cfg.get('相片备注')
           == '{{项目类型}}{{标段}}{{乡（镇）}}{{村（社区）}}{{小班号}}小班'
           and cfg.get('相片文件名') == '{{标段}}_{{小班号}}_{{时间}}'
           and str(cfg.get('压缩最长边')) == '1440'
           and abs(float(cfg.get('压缩质量')) - 0.8) < 1e-9,
           str({k: cfg.get(k) for k in
                ('目录', '相片备注', '相片文件名', '压缩最长边', '压缩质量')})[:160])
        ridx = next(i for i, x in enumerate(rows)
                    if str(dict(zip(hdrs, x)).get('小班号', '')).strip() == 'XP-1320')

        f = cli.get(f'/api/workbooks/{wid}/flow?row={ridx}', headers=H).get_json()
        steps = f.get('steps') or []
        ok('T2 默认流程 = 路由第一条「整地」（8 个环节）',
           f.get('flow_sheet') == ROUTES[0] and len(steps) == 8,
           f"{f.get('flow_sheet')} {len(steps)} 环节")
        ok('T2b flow_list 按路由顺序给全部 5 条',
           [x.get('sheet') for x in (f.get('flow_list') or [])] == ROUTES,
           str([x.get('sheet') for x in (f.get('flow_list') or [])]))
        ok('T2c 设计值表头映射成 std（面积/密度按小班取值）',
           f.get('std', {}).get('核实施工面积') == 112.52
           and f.get('std', {}).get('核实密度') == 54,
           str({k: f.get('std', {}).get(k) for k in ('核实施工面积', '核实密度')}))
        ctl = next(c for s in steps for c in s['controls'] if c['name'] == '核实施工面积')
        ok('T2d 控件带出「设计值表头」（界面据此放大显示设计值）',
           ctl.get('design') == '设计面积（亩）', str(ctl.get('design')))
        jc = next(c for s in steps for c in s['controls'] if c['name'] == '核实总株数')
        ok('T2i 模板 `数值(株→≥设计值)` → 控件带出已解析的 judge_* 字段',
           jc.get('judge') == '≥设计值' and jc.get('judge_op') == '≥'
           and jc.get('judge_ref') == '设计值' and jc.get('unit') == '株',
           str({k: jc.get(k) for k in ('judge', 'judge_op', 'judge_ref', 'unit')}))
        ok('T2e 备注行被认成流程共用备注（note_field=备注，隐藏）',
           f.get('note_field') == '备注' and f.get('note_step_no') == 7,
           f"{f.get('note_field')} / {f.get('note_step_no')}")
        step1 = next(s for s in steps if s['step'] == '核实施工面积')
        inline = next((c for c in step1['controls'] if c.get('note_for')), {})
        ok('T2h 选项带 →备注 时就地生成隐藏文本框（同环节，条件=不正确）',
           inline.get('name') == '核实施工面积说明' and inline.get('hidden') is True
           and inline.get('label') == '备注'
           and inline.get('show_on') == {'核实施工面积': ['不正确']},
           str(inline)[:140])
        ok('T2f 整地可填、回塘底肥锁定（前置未完成）',
           f.get('flow_locked') is False
           and (f.get('flow_list') or [{}])[1].get('available') is False,
           str([(x.get('sheet'), x.get('available')) for x in f.get('flow_list') or []]))
        patrol = (f.get('flow_list') or [{}])[-1]
        ok('T2g 独立流程「日常巡护」排最后：无前置、可多次、不抢默认',
           patrol.get('sheet') == ROUTES[4] and patrol.get('independent') is True
           and patrol.get('available') is True and patrol.get('repeat') is True
           and f.get('flow_sheet') != ROUTES[4],
           f"{patrol.get('sheet')} independent={patrol.get('independent')} "
           f"available={patrol.get('available')} default={f.get('flow_sheet')}")

        # ── v0.42 附件上传：模板「流程-日常巡护」第 1 环节要求 3 份（图片≥1 / PDF≥1）──
        fp = cli.get(f'/api/workbooks/{wid}/flow?row={ridx}&flow={ROUTES[4]}',
                     headers=H).get_json()
        spec = ((fp.get('steps') or [{}])[0].get('attachments') or [{}])[0]
        ok('T11a 模板附件列解析：上传「巡护资料」≥3 份，图片≥1 / PDF≥1',
           spec.get('kind') == 'upload' and spec.get('label') == '巡护资料'
           and spec.get('min') == 3
           and (spec.get('types') or {}) == {'image': 1, 'pdf': 1},
           str(spec)[:140])
        r = try_att(cli, H, wid, ridx, ROUTES[4], 1, '不存在的标签', 'a.png', PNG)
        ok('T11b 上传到未声明的标签 → 400（服务端不认客户端随便传）',
           r.status_code == 400, str((r.get_json() or {}).get('error'))[:80])
        r = try_att(cli, H, wid, ridx, ROUTES[4], 1, '巡护资料', 'a.txt', b'hello')
        ok('T11c 非图片/PDF（txt）→ 400（按魔数拒绝）',
           r.status_code == 400, str((r.get_json() or {}).get('error'))[:80])

        three_img = att_entry(cli, H, wid, ridx, ROUTES[4], 1, '巡护资料',
                              [('p1.png', PNG), ('p2.jpg', JPG), ('p3.png', PNG)])
        r = batch(cli, H, wid, ridx, ROUTES[4],
                  full_steps(PATROL_PASS, {1: [three_img]}))
        err = str((r.get_json() or {}).get('error', ''))
        ok('T11d 3 张图片、没有 PDF → 400（类型下限，数量够也不行）',
           r.status_code == 400 and 'PDF' in err, err[:100])

        three_pdf = att_entry(cli, H, wid, ridx, ROUTES[4], 1, '巡护资料',
                              [('d1.pdf', PDF), ('d2.pdf', PDF), ('d3.pdf', PDF)])
        r = batch(cli, H, wid, ridx, ROUTES[4],
                  full_steps(PATROL_PASS, {1: [three_pdf]}))
        err = str((r.get_json() or {}).get('error', ''))
        ok('T11e 3 个 PDF、没有图片 → 400（图片下限）',
           r.status_code == 400 and '图片' in err, err[:100])

        patrol_att1 = att_entry(cli, H, wid, ridx, ROUTES[4], 1, '巡护资料',
                                [('巡护1.png', PNG), ('巡护2.jpg', JPG),
                                 ('巡护记录.pdf', PDF)])
        two = dict(patrol_att1, ids=patrol_att1['ids'][:2],
                   names=patrol_att1['names'][:2], types=patrol_att1['types'][:2],
                   count=2)
        r = batch(cli, H, wid, ridx, ROUTES[4],
                  full_steps(PATROL_PASS, {1: [two]}))
        err = str((r.get_json() or {}).get('error', ''))
        ok('T11f 只认领 2 份 → 400 且报「需 3 份」',
           r.status_code == 400 and '需 3 份' in err, err[:100])

        # ── 老接口不许走 ──
        r = cli.post(f'/api/workbooks/{wid}/flow', headers=H, json={
            'row': ridx, 'step_no': 1, 'values': {'核实施工面积': '正确'}})
        ok('T10 v2 工作簿走老 /flow 提交 → 400（必须整条提交）',
           r.status_code == 400 and 'batch' in str((r.get_json() or {}).get('error', '')),
           str((r.get_json() or {}).get('error'))[:80])

        # ── 负例（不写库）──
        r = batch(cli, H, wid, ridx, ROUTES[0], full_steps({1: {'核实施工面积': '正确'}}))
        err = str((r.get_json() or {}).get('error', ''))
        ok('T5 整地只填一项 → 400 且列出缺项', r.status_code == 400 and '核实密度' in err,
           err[:80])

        r = batch(cli, H, wid, ridx, ROUTES[0], full_steps(
            {**LAND_PASS, 6: {'核实总株数': -5}}))
        err = str((r.get_json() or {}).get('error', ''))
        ok('T5b 核实总株数 -5 → 400（数值规则 ≥0）',
           r.status_code == 400 and '核实总株数' in err, err[:80])

        r = batch(cli, H, wid, ridx, ROUTES[0], full_steps(
            {**LAND_PASS, 8: {'是否通过': '不通过'}}))
        ok('T4 结论不通过但没写备注 → 400（≥2 个汉字）',
           r.status_code == 400 and '备注' in str((r.get_json() or {}).get('error', '')),
           str((r.get_json() or {}).get('error'))[:80])

        r = batch(cli, H, wid, ridx, ROUTES[0], full_steps(
            {**LAND_PASS, 1: {'核实施工面积': '不正确'}}))
        ok('T4b 选了「不正确」（→备注）没写备注 → 400',
           r.status_code == 400 and '备注' in str((r.get_json() or {}).get('error', '')),
           str((r.get_json() or {}).get('error'))[:80])

        # ── 独立流程：先交巡护，不许影响主线 ──
        r = batch(cli, H, wid, ridx, ROUTES[4],
                  full_steps(PATROL_PASS, {1: [patrol_att1]}),
                  geo={'lat': 22.1234, 'lon': 102.1234, 'acc': 8})
        ok('T9a 独立流程可先于主线提交（rec_no=1，附件一并认领）',
           r.status_code == 200 and (r.get_json() or {}).get('rec_no') == 1,
           f'{r.status_code} {(r.get_json() or {}).get("rec_no")} '
           f'{str((r.get_json() or {}).get("error"))[:60]}')

        aid = patrol_att1['ids'][0]
        r = cli.get(f'/api/attachments/{aid}', headers=H)
        ok('T12a 附件可下载：内容与上传一致（登录鉴权）',
           r.status_code == 200 and r.data == PNG,
           f'{r.status_code} {len(r.data)}B')
        r = app.test_client().get(f'/api/attachments/{aid}')
        ok('T12b 未登录取附件 → 401（下载不是公开链接）',
           r.status_code == 401, str(r.status_code))
        con = sqlite3.connect(_M.DB_PATH)
        bound = con.execute('SELECT COUNT(*) FROM flow_att WHERE workbook_id=?'
                            ' AND step_no=1 AND rec_no=1', (wid,)).fetchone()[0]
        con.close()
        ok('T12c 提交后 flow_att 绑定 rec_no=1（3 个文件）', bound == 3, f'{bound} 个')

        r = batch(cli, H, wid, ridx, ROUTES[4],
                  full_steps(PATROL_PASS, {1: [patrol_att1]}), rec_no=2)
        err = str((r.get_json() or {}).get('error', ''))
        ok('T12d 同一批附件再报给另一次（rec_no=2）→ 400（不许一份文件反复用）',
           r.status_code == 400 and '别的记录' in err, err[:100])

        f = cli.get(f'/api/workbooks/{wid}/flow?row={ridx}', headers=H).get_json()
        ok('T9b 巡护提交后主线默认仍是整地、回塘仍锁定（互不阻塞）',
           f.get('flow_sheet') == ROUTES[0] and f.get('flow_locked') is False
           and (f.get('flow_list') or [{}])[1].get('available') is False,
           f"default={f.get('flow_sheet')} "
           f"回塘locked={not (f.get('flow_list') or [{}])[1].get('available')}")

        r = batch(cli, H, wid, ridx, ROUTES[1], full_steps(BASE_PASS))
        ok('T3 前置未完成提交回塘底肥 → 409（流程路由闸门）',
           r.status_code == 409 and '整地' in str((r.get_json() or {}).get('error', '')),
           str((r.get_json() or {}).get('error'))[:80])

        # ── v0.45 通过判断：实测低于设计值时不许「通过」──
        r = batch(cli, H, wid, ridx, ROUTES[0], full_steps(
            {**LAND_PASS, 6: {'核实总株数': 6000}}))
        err = str((r.get_json() or {}).get('error', ''))
        ok('T5c 株数 6000 < 设计 6076 却选「通过」→ 400（通过判断拦截）',
           r.status_code == 400 and '通过判断' in err and '6000' in err, err[:120])

        design = f.get('std', {}).get('核实总株数')
        ok('T5d 设计值从当前小班取（XP-1320 总株数设计值可查）',
           design is not None and str(design).strip() != '', f'设计={design}')

        # ── 主线：整地 → 回塘底肥 → 栽植 → 完工 ──
        r = batch(cli, H, wid, ridx, ROUTES[0], full_steps(LAND_PASS))
        ok('T6 整地整条提交 → 200，rec_no=0（一次性流程）',
           r.status_code == 200 and (r.get_json() or {}).get('rec_no') == 0,
           f'{r.status_code} {str((r.get_json() or {}).get("error"))[:80]}')

        f = cli.get(f'/api/workbooks/{wid}/flow?row={ridx}', headers=H).get_json()
        ok('T7 整地落库为通过，下一条「回塘底肥」可填',
           f.get('flow_sheet') == ROUTES[1] and f.get('flow_locked') is False
           and (f.get('flow_list') or [{}])[0].get('done') is True,
           f"{f.get('flow_sheet')} locked={f.get('flow_locked')}")
        fl0 = cli.get(f'/api/workbooks/{wid}/flow?row={ridx}&flow={ROUTES[0]}',
                      headers=H).get_json()
        ok('T7b 整地的填报值原样回显（含设计值对照）',
           str((fl0.get('values') or {}).get('1', {}).get('核实施工面积')) == '正确'
           and fl0.get('std', {}).get('核实施工面积') == 112.52,
           str((fl0.get('values') or {}).get('1'))[:80])
        ok('T7c 常量设计值也进 std（回塘底肥 1.5kg/株）',
           f.get('std', {}).get('有机肥1.5kg/株') == '1.5kg/株',
           str(f.get('std', {}).get('有机肥1.5kg/株')))
        lg = fl0.get('flow_log') or []
        ok('T7g 主线流程也出日志：结论/操作人/时间/环节数',
           len(lg) == 1 and lg[0].get('result') == '通过'
           and lg[0].get('operator') == ADMIN[0] and lg[0].get('at')
           and lg[0].get('steps') == 8,
           str(lg[:1])[:120])

        r = batch(cli, H, wid, ridx, ROUTES[1], full_steps(
            {**BASE_PASS, 2: {'有机肥1.5kg/株': '未实施施肥'}}))
        err = str((r.get_json() or {}).get('error', ''))
        ok('T4c 多个错误项（有机肥未实施施肥）没写就地备注 → 400',
           r.status_code == 400 and '备注' in err and '有机肥' in err, err[:100])

        r = batch(cli, H, wid, ridx, ROUTES[1], full_steps(
            {**BASE_PASS, 2: {'有机肥1.5kg/株': '未实施施肥',
                              '有机肥1.5kg/株说明': '现场确认尚未施肥'}}))
        ok('T7d 回塘底肥提交成功（错误项就地备注已填）', r.status_code == 200,
           f'{r.status_code} {str((r.get_json() or {}).get("error"))[:80]}')

        r = batch(cli, H, wid, ridx, ROUTES[2], full_steps(PLANT_PASS))
        ok('T7e 栽植撑杆覆膜提交成功', r.status_code == 200,
           f'{r.status_code} {str((r.get_json() or {}).get("error"))[:80]}')

        r = batch(cli, H, wid, ridx, ROUTES[3], full_steps(
            {**DONE_PASS, 8: {'栽植成活率': 80}}))
        err = str((r.get_json() or {}).get('error', ''))
        ok('T5e 成活率 80% < 85% 选「通过」→ 400（百分数规则）',
           r.status_code == 400 and '通过判断' in err and '85' in err, err[:120])

        r = batch(cli, H, wid, ridx, ROUTES[3], full_steps(
            {**DONE_PASS, 13: {'样地调查成活株数': 60},
             14: {'样地调查死亡株数': 40}}))
        err = str((r.get_json() or {}).get('error', ''))
        ok('T5f 自动公式算出 60% < 85% 也拦（自动值同规则判定）',
           r.status_code == 400 and '通过判断' in err, err[:120])

        r = batch(cli, H, wid, ridx, ROUTES[3], full_steps(DONE_PASS))
        ok('T7f 完工验收提交成功', r.status_code == 200,
           f'{r.status_code} {str((r.get_json() or {}).get("error"))[:80]}')

        f = cli.get(f'/api/workbooks/{wid}/flow?row={ridx}&flow={ROUTES[3]}', headers=H).get_json()
        got = str((f.get('values') or {}).get('15', {}).get('样地调查成活率'))
        ok('T8 自动公式由服务端算：90/(90+10)×100 = 90',
           got in ('90', '90.0'), f'落库={got}')
        ok('T8b 主线四条流程全部 done',
           all(x.get('done') for x in (f.get('flow_list') or [])[:4]),
           str([(x.get('sheet'), x.get('done')) for x in f.get('flow_list') or []]))
        f = cli.get(f'/api/workbooks/{wid}/flow?row={ridx}', headers=H).get_json()
        ok('T8c 主线全完成后默认仍停在完工验收，不跳到独立巡护',
           f.get('flow_sheet') == ROUTES[3] and f.get('flow_independent') is False,
           f"{f.get('flow_sheet')} independent={f.get('flow_independent')}")

        # ── 现场特殊：判断不过也允许「不通过 + 备注」，不把人堵死 ──
        ridx2 = next(i for i in range(len(rows)) if i != ridx)
        r = batch(cli, H, wid, ridx2, ROUTES[0], full_steps(
            {**LAND_PASS, 6: {'核实总株数': 100}, 7: {'备注': '现场核实株数不足'},
             8: {'是否通过': '不通过'}}))
        ok('T5g 判断不过仍可选「不通过 + 备注」提交（留痕，不当场拦死）',
           r.status_code == 200,
           f'{r.status_code} {str((r.get_json() or {}).get("error"))[:100]}')

        # ── 可重复 / 独立流程：日常巡护 ──
        att2 = att_entry(cli, H, wid, ridx, ROUTES[4], 1, '巡护资料',
                         [('巡护3.png', PNG), ('巡护4.jpg', JPG), ('记录2.pdf', PDF)])
        att3 = att_entry(cli, H, wid, ridx, ROUTES[4], 1, '巡护资料',
                         [('巡护5.png', PNG), ('巡护6.jpg', JPG), ('记录3.pdf', PDF)])
        r1 = batch(cli, H, wid, ridx, ROUTES[4],
                   full_steps(PATROL_PASS, {1: [att2]}))
        r2 = batch(cli, H, wid, ridx, ROUTES[4],
                   full_steps(PATROL_PASS, {1: [att3]}))
        ok('T9c 独立流程连交两次 → rec_no 2/3（接着先前那次，各带新附件）',
           r1.status_code == 200 and r2.status_code == 200
           and (r1.get_json() or {}).get('rec_no') == 2
           and (r2.get_json() or {}).get('rec_no') == 3,
           f'{r1.status_code}/{r2.status_code} '
           f'{(r1.get_json() or {}).get("rec_no")}/{(r2.get_json() or {}).get("rec_no")}')

        r3 = batch(cli, H, wid, ridx, ROUTES[4],
                   full_steps(PATROL_PASS, {1: [patrol_att1]}), rec_no=1)
        fp = cli.get(f'/api/workbooks/{wid}/flow?row={ridx}&flow={ROUTES[4]}',
                     headers=H).get_json()
        m = (fp.get('multi') or {}).get('1') or {}
        ok('T9d 带 rec_no=1 提交 = 改那一次，不新增（仍 3 次，附件可沿用原记录）',
           r3.status_code == 200 and (r3.get_json() or {}).get('rec_no') == 1 and len(m) == 3,
           f'{(r3.get_json() or {}).get("rec_no")} / {len(m)} 条')
        ok('T9e 巡护多选落库为数组',
           '林地清理、整地巡护' in str((m.get('1') or {}).get('values', {}).get('巡护情况')),
           str((m.get('1') or {}).get('values', {}).get('巡护情况'))[:60])
        got_ids = ((m.get('1') or {}).get('att') or [{}])[0].get('ids') or []
        ok('T12e GET 回显已提交附件（3 个 id 与上传一致）',
           got_ids == patrol_att1['ids'], str(got_ids))
        ok('T9f 巡护在 GET 里标记 independent / repeat / 定期',
           fp.get('flow_independent') is True and fp.get('flow_repeat') is True
           and fp.get('flow_trigger') == '定期',
           f"{fp.get('flow_independent')} {fp.get('flow_repeat')} {fp.get('flow_trigger')}")

        # ── v0.40 提交日志：每次通过/不通过都留痕（含改那一次）──
        lg = fp.get('flow_log') or []
        ok('T9g 独立流程日志 4 条（含改第 1 次），结论/操作人/IP 齐全且新的在前',
           len(lg) == 4
           and all(x.get('result') == '通过' and x.get('operator') == ADMIN[0]
                   and x.get('ip') for x in lg)
           and [x.get('rec_no') for x in lg] == [1, 3, 2, 1],
           f'{len(lg)} 条 rec={[x.get("rec_no") for x in lg]}')
        g0 = (lg[-1] or {}).get('geo') or {}
        ok('T9h 日志记下提交瞬间经纬度（22.1234 / 102.1234）',
           str(g0.get('lat')) == '22.1234' and str(g0.get('lon')) == '102.1234'
           and str(g0.get('acc')) == '8',
           str(g0)[:80])

        # ── v0.48 导出：数据 sheet 原样 + 追加「流程填报」长表 ──
        buf, _name, _mime = _M._export_workbook(wid)
        ew = openpyxl.load_workbook(buf)
        ok('T13 导出追加「流程填报」，原数据 sheet 保留',
           '流程填报' in ew.sheetnames and '设计值' in ew.sheetnames,
           str(ew.sheetnames)[:100])
        ws2 = ew['流程填报']
        vals = [list(r) for r in ws2.iter_rows(min_row=2, values_only=True)]
        ok('T13b 明细含小班/流程/结论/操作人/附件',
           any(str(r[0]) == 'XP-1320' and r[1] == ROUTES[4] and r[6] == '通过'
               and r[7] == ADMIN[0] for r in vals)
           and any('：' in str(r[10]) and '.png' in str(r[10]) for r in vals),
           f'{len(vals)} 行 {vals[:1]}'[:160])
        ok('T13c 数据 sheet 行数未被流程明细改变（149 行 + 表头）',
           ew['设计值'].max_row == 150, str(ew['设计值'].max_row))

        # ── v0.48 后台填报进度：矩阵 + 明细 ──
        pg = cli.get(f'/admin/api/progress?wid={wid}', headers=H).get_json()
        flows_pg = pg.get('flows') or []
        ok('T14 后台进度：5 条流程且独立巡护排最后',
           [f['sheet'] for f in flows_pg] == ROUTES,
           str([f['sheet'] for f in flows_pg]))
        tot = (pg.get('totals') or {}).get(ROUTES[4]) or {}
        ok('T14b 巡护统计：通过 ≥1、未填 + 通过 = 149、提交次数 ≥1',
           tot.get('pass', 0) >= 1 and tot.get('pass', 0) + tot.get('none', 0) == 149
           and tot.get('submissions', 0) >= 1, str(tot))
        cell = next((c for r in (pg.get('rows') or []) if r['idx'] == ridx
                     for c in (r.get('cells') or []) if c['flow'] == ROUTES[4]), {})
        ok('T14c 进度单元格带最近操作人/时间与次数',
           cell.get('state') == 'pass' and cell.get('operator') == ADMIN[0]
           and cell.get('count', 0) >= 1, str(cell)[:120])
        de = cli.get(f'/admin/api/progress/detail?wid={wid}&row={ridx}&flow={ROUTES[4]}',
                     headers=H).get_json()
        dvals = [v for x in (de.get('rows') or []) for s in x['steps'] for v in s['values']]
        ok('T14d 明细含巡护情况与附件清单',
           any('林地清理' in v['value'] for v in dvals)
           and any('.png' in s.get('att', '') for x in (de.get('rows') or [])
                   for s in x['steps']),
           f"{len(dvals)} 个值")
        ok('T14e 明细带提交日志（结论/操作人）',
           any(a.get('result') == '通过' and a.get('operator') == ADMIN[0]
               for a in (de.get('audits') or [])), str(len(de.get('audits') or [])))
    finally:
        con = sqlite3.connect(_M.DB_PATH)
        con.execute('DELETE FROM workbooks WHERE id=?', (wid,))
        con.execute('DELETE FROM flow_records WHERE workbook_id=?', (wid,))
        con.execute('DELETE FROM flow_audit WHERE workbook_id=?', (wid,))
        con.execute('DELETE FROM flow_att WHERE workbook_id=?', (wid,))
        con.commit()
        con.close()
        shutil.rmtree(_M.UPLOAD_DIR / str(wid), ignore_errors=True)
        shutil.rmtree(_M.ATT_DIR / str(wid), ignore_errors=True)

    print('\n' + '=' * 62)
    print(f'通过 {_pass} / 失败 {_fail}')
    return 0 if _fail == 0 else 1


if __name__ == '__main__':
    sys.exit(main())
