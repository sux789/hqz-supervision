#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""参数解析与唯一键的单元测试（v0.23，C11）。

用法：python tests/test_unit.py      # 全绿退出码 0

只测纯函数，不碰数据库与网络；涉及真实模板/HTTP 的部分在 tests/test_e2e.py。
"""
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

from param_parser import (  # noqa: E402
    ParamError, check_unique_column, export_filters_of, filter_options,
    filter_rows, log_fields_of, merge_rows, parse_params, percent_cols_of,
    row_key_column, unique_key_of)
from main import _norm_cell, _pick_data_sheet  # noqa: E402

HDR = ['key', 'value', '类型', '默认值', '说明', '示例']
HEADERS = ['项目类型', '小班号', '标段', '验收人', '验收日期', '验收结果', '验收备注',
           '株数强度', '混交比例']

_pass = _fail = 0


def ok(name, cond, detail=''):
    global _pass, _fail
    _pass, _fail = _pass + bool(cond), _fail + (not cond)
    print(f'  {"✅" if cond else "❌"} {name}  {detail}')


def prows(*rows):
    """构造参数 sheet 行：[表头行] + 数据行（自动补足 6 列）。"""
    return [(1, HDR)] + [(i + 2, list(r) + [''] * (6 - len(r))) for i, r in enumerate(rows)]


def raises(name, fn, must_contain=''):
    try:
        fn()
        ok(name, False, '未报错（预期应报错）')
    except ParamError as e:
        ok(name, must_contain in str(e) if must_contain else True, str(e))


def passes(name, fn, check=None):
    try:
        r = fn()
    except ParamError as e:
        ok(name, False, str(e))
        return None
    ok(name, check(r) if check else True, repr(r))
    return r


# ── 参数三件套（A）──────────────────────────────────────────
print('=== A. 参数解析：必填与按需必填 ===')
FULL = [['可编辑列', '验收人;验收日期'], ['功能', '轨迹'],
        ['相片备注', '{{小班号}}'], ['相片文件名', '{{小班号}}_{{时间}}'],
        ['目录', '{{项目类型}}/{{小班号}}']]

passes('A1 三件套齐全 → 通过', lambda: parse_params(prows(*FULL), HEADERS))
raises('A2 缺「功能」→ 报错（恒必填）', lambda: parse_params(prows(*FULL[0:1]), HEADERS), '功能')
raises('A3 缺「可编辑列」→ 报错（恒必填）',
       lambda: parse_params(prows(*[[r for r in FULL if r[0] != '可编辑列']][0]), HEADERS), '可编辑列')
passes('A4 不拍照（功能=轨迹）且无三件套 → 通过（v0.23 起按需必填）',
       lambda: parse_params(prows(['可编辑列', '验收人'], ['功能', '轨迹']), HEADERS))
raises('A5 功能含拍照但无三件套 → 报错',
       lambda: parse_params(prows(['可编辑列', '验收人'], ['功能', '拍照']), HEADERS), '相片备注')
raises('A6 功能含视频但无三件套 → 报错',
       lambda: parse_params(prows(['可编辑列', '验收人'], ['功能', '视频']), HEADERS), '相片备注')
passes('A7 功能含拍照且三件套齐全 → 通过',
       lambda: parse_params(prows(['可编辑列', '验收人'], ['功能', '拍照;轨迹'],
                                  ['相片备注', '{{小班号}}'], ['相片文件名', '{{小班号}}'],
                                  ['目录', '{{项目类型}}']), HEADERS))
raises('A8 未知 key → 报错', lambda: parse_params(prows(*FULL, ['唯一键', '小班号']), HEADERS), '未知 key')
raises('A9 占位符不是表头 → 报错',
       lambda: parse_params(prows(*FULL, ['相片备注', '{{地块号}}']), HEADERS), '不是数据 sheet 表头')

# ── 唯一键（B）────────────────────────────────────────────
print('\n=== B. unique-key 声明与校验 ===')
UK = ['unique-key', '小班号']
passes('B1 unique-key=表头列 → 通过', lambda: parse_params(prows(*FULL, UK), HEADERS),
       lambda c: c['unique-key'] == '小班号')
raises('B2 指向不存在的列 → 报错',
       lambda: parse_params(prows(*FULL, ['unique-key', '地块号']), HEADERS), '不在数据 sheet 表头')
raises('B3 多值 → 报错',
       lambda: parse_params(prows(*FULL, ['unique-key', '小班号;标段']), HEADERS), '不支持多值')
raises('B4 出现在「可编辑列」里 → 报错',
       lambda: parse_params(prows(['可编辑列', '小班号;验收人'], ['功能', '轨迹'], UK), HEADERS), '不能出现在')
passes('B5 未声明 → 不校验（即使数据重复）', lambda: check_unique_column(
    [['T', 'X'], ['T', 'X']], HEADERS, {}), lambda r: r is None)
ok('B6 唯一键重复 → 返回带行号错误串',
   isinstance(check_unique_column([['T', 'X'], ['T', 'Y'], ['T', 'X']], HEADERS,
                                 {'unique-key': '小班号'}), str))
ok('B7 唯一键为空 → 返回带行号错误串',
   isinstance(check_unique_column([['T', ''], ['T', 'Y']], HEADERS,
                                 {'unique-key': '小班号'}), str))
ok('B8 全唯一 → None',
   check_unique_column([['T', 'X'], ['T', 'Y']], HEADERS, {'unique-key': '小班号'}) is None)
ok('B9 唯一键列不在表头 → 返回错误串',
   isinstance(check_unique_column([['T', 'X']], HEADERS, {'unique-key': '地块号'}), str))

# ── 生效值助手（C）────────────────────────────────────────
print('\n=== C. 生效值助手：声明优先 / 缺省回落 ===')
passes('C1 unique_key_of 未声明 → 空串', lambda: unique_key_of({}), lambda r: r == '')
passes('C2 row_key_column 未声明 → 回落小班号', lambda: row_key_column({}), lambda r: r == '小班号')
passes('C3 row_key_column 声明优先', lambda: row_key_column({'unique-key': '地块号'}),
       lambda r: r == '地块号')
passes('C4 log_fields_of 缺省 = 内置里真实存在的列',
       lambda: log_fields_of(HEADERS, {}),
       lambda r: r == ['验收人', '验收日期', '验收结果', '验收备注'])
passes('C5 log_fields_of 声明优先',
       lambda: log_fields_of(HEADERS, {'日志字段': '验收人;验收备注'}),
       lambda r: r == ['验收人', '验收备注'])
passes('C6 percent_cols_of 缺省 = 列名含「强度」', lambda: percent_cols_of(HEADERS, {}),
       lambda r: r == {7})
passes('C7 percent_cols_of 声明优先（可覆盖非强度列）',
       lambda: percent_cols_of(HEADERS, {'百分比列': '混交比例'}), lambda r: r == {8})

# ── 数据 sheet 定位（D）───────────────────────────────────
print('\n=== D. 数据 sheet 定位（C11）===')


def pick(sheets):
    try:
        return _pick_data_sheet(sheets)
    except ParamError as e:
        return f'ERR:{e}'


ok('D1 有 data → 取 data', pick(['data', '参数']) == 'data', pick(['data', '参数']))
ok('D2 有 data + 说明 → 仍取 data（多余页可容忍）',
   pick(['data', '说明', '参数']) == 'data', pick(['data', '说明', '参数']))
ok('D3 无 data、单非参数 sheet → 取它（旧模板兼容）',
   pick(['中幼林抚育验收', '参数']) == '中幼林抚育验收')
ok('D4 无 data、多非参数 sheet → 报错', pick(['甲', '乙', '参数']).startswith('ERR'))
ok('D5 只有参数 → 报错', pick(['参数']).startswith('ERR'))
ok('D6 歧义报错里列出全部 sheet 名', '甲' in pick(['甲', '乙', '参数']) and '乙' in pick(['甲', '乙', '参数']))

# ── 导出筛选（E，v0.25）──────────────────────────────────────
print('\n=== E. 导出筛选：声明、候选值、筛行 ===')
FH = ['标段', '乡镇', '验收人', '验收日期']
FROWS = [['10标段', 'A乡', '张三', '2026-09-10'],
         ['6标段', 'B乡', '李四', '2026-09-11 09:30:00'],
         ['7标段', 'A乡', '张三', '2026-09-12'],
         ['8标段', '', '', '']]
FCFG = {'导出筛选': '标段|select;乡镇|select;验收人|select;验收日期|date'}


def fsel(filters):
    return filter_rows(FROWS, FH, filters, dict(export_filters_of(FH, FCFG)))


passes('E1 声明字段 = 声明 ∩ 表头（顺序保持）',
       lambda: export_filters_of(FH, FCFG),
       lambda r: r == [('标段', 'select'), ('乡镇', 'select'), ('验收人', 'select'), ('验收日期', 'date')])
passes('E2 表头里没有的声明字段被剔除（不提示筛选）',
       lambda: export_filters_of(FH, {'导出筛选': '标段|select;不存在的列|select;也要剔除|date'}),
       lambda r: r == [('标段', 'select')])
passes('E3 未声明「导出筛选」→ 空清单（不显示弹框）',
       lambda: export_filters_of(FH, {}), lambda r: r == [])
passes('E4 下拉候选值按「数据中首次出现」去重（不按字典序，中文值不会被打乱）',
       lambda: filter_options(FROWS, FH, '标段'),
       lambda r: r == ['10标段', '6标段', '7标段', '8标段'])
passes('E5 候选值跳过空单元格', lambda: filter_options(FROWS, FH, '乡镇'),
       lambda r: r == ['A乡', 'B乡'])
passes('E6 字段不在表头 → 无候选值', lambda: filter_options(FROWS, FH, '小班号'),
       lambda r: r == [])
passes('E7 select 精确匹配', lambda: [r[0] for r in fsel({'标段': '7标段'})], lambda r: r == ['7标段'])
passes('E8 date 单日匹配（值带时间也命中）',
       lambda: [r[0] for r in fsel({'验收日期': '2026-09-11'})], lambda r: r == ['6标段'])
passes('E9 多条件 AND', lambda: len(fsel({'标段': '7标段', '乡镇': 'A乡'})), lambda r: r == 1)
passes('E10 多条件之一不匹配 → 0 行', lambda: len(fsel({'标段': '7标段', '乡镇': 'B乡'})),
       lambda r: r == 0)
passes('E11 空值条件不生效', lambda: len(fsel({'标段': '', '验收人': None})), lambda r: r == 4)
passes('E12 未声明的字段被忽略（不会把数据筛空）',
       lambda: len(filter_rows(FROWS, FH, {'小班号': 'X'}, {})), lambda r: r == 4)
passes('E13 声明字段值不存在 → 0 行', lambda: len(fsel({'标段': '不存在'})), lambda r: r == 0)

# E14~E16 用一组与 FH 匹配的合法基底（功能不含拍照，故拍照三件套可不写）
FBASE = [['可编辑列', '验收人;验收日期'], ['功能', '轨迹']]

raises('E14 控件类型不支持 → 报错',
       lambda: parse_params(prows(*FBASE, ['导出筛选', '标段|checkbox']), FH), '不支持')
raises('E15 字段不在表头 → 报错',
       lambda: parse_params(prows(*FBASE, ['导出筛选', '小班号|select']), FH), '不在数据 sheet 表头')
raises('E16 缺控件类型（只有字段名）→ 报错',
       lambda: parse_params(prows(*FBASE, ['导出筛选', '标段']), FH), '格式应为')
passes('E17 合法声明 → 通过',
       lambda: parse_params(prows(*FBASE, ['导出筛选', '标段|select;验收日期|date']), FH),
       lambda c: c['导出筛选'] == '标段|select;验收日期|date')

# ── merge_rows：按唯一键合并更新（F，v0.26）───────────────────
print('\n=== F. 按唯一键合并更新（merge_rows）===')

OH = ['小班号', '小班面积', '验收人', '验收备注']
NH = ['小班号', '验收人', '验收备注', '小班面积', '新增列']     # 换序 + 加列
OROWS = [['A1', 100.5, '雷华雄', '甲方已确认'],
         ['A2', 200, '', '待复核'],
         ['A3', 300, '张三', '仅旧表有']]
NROWS = [['A2', '李四', '', 999, 'N2'],
         ['A1', '不该覆盖', '新表备注', 111, 'N1'],
         ['A4', '', '', 444, 'N4']]
PRESERVE = ['验收人', '验收备注']


def mrg(oh, orows, nh, nrows, key='小班号', preserve=PRESERVE):
    return merge_rows(oh, orows, nh, nrows, key, list(preserve))


passes('F1 匹配行原地合并；输出列序 = 新表列序',
       lambda: mrg(OH, OROWS, NH, NROWS)[0],
       lambda r: r == [['A1', '雷华雄', '甲方已确认', 111, 'N1'],
                       ['A2', '李四', '待复核', 999, 'N2'],
                       ['A3', '张三', '仅旧表有', 300, ''],
                       ['A4', '', '', 444, 'N4']])
passes('F2 老行顺序不变（row_idx 不错位，日志才不会串行）',
       lambda: [r[0] for r in mrg(OH, OROWS, NH, NROWS)[0]],
       lambda r: r == ['A1', 'A2', 'A3', 'A4'])
passes('F3 统计：匹配 2 / 新增 1 / 仅旧 1 / 冲突 2',
       lambda: {k: mrg(OH, OROWS, NH, NROWS)[1][k]
                for k in ('matched', 'added', 'kept_old_only', 'conflicts')},
       lambda r: r == {'matched': 2, 'added': 1, 'kept_old_only': 1, 'conflicts': 2})
passes('F4 可编辑列：旧值非空优先（新表覆盖不了人工填写）',
       lambda: mrg(OH, OROWS, NH, NROWS)[0][0][2], lambda r: r == '甲方已确认')
passes('F5 可编辑列：旧值为空 → 取新表的值',
       lambda: mrg(OH, OROWS, NH, NROWS)[0][1][1], lambda r: r == '李四')
passes('F6 非可编辑列：以新表为准（改小班面积能生效）',
       lambda: mrg(OH, OROWS, NH, NROWS)[0][0][3], lambda r: r == 111)
passes('F7 数值类型保留（不会变成字符串）',
       lambda: [type(mrg(OH, OROWS, NH, NROWS)[0][0][3]).__name__,
                type(mrg(OH, OROWS, NH, NROWS)[0][2][3]).__name__],
       lambda r: r == ['int', 'int'])
passes('F8 新表新增列：老行取新表值、仅旧数据的行留空',
       lambda: [mrg(OH, OROWS, NH, NROWS)[0][0][4], mrg(OH, OROWS, NH, NROWS)[0][2][4]],
       lambda r: r == ['N1', ''])

OH2 = ['小班号', '小班面积', '验收人']
passes('F9 非保护列被新表空值清空 → blanked_nonempty 计数（要在弹框警告）',
       lambda: (mrg(OH2, [['B1', 108.87, '雷华雄']], OH2, [['B1', '', '新名']])[0][0],
                mrg(OH2, [['B1', 108.87, '雷华雄']], OH2, [['B1', '', '新名']])[1]['blanked_nonempty']),
       lambda r: r == (['B1', '', '雷华雄'], 1))
passes('F10 新表删列 → 记录会丢失的有值格数与列名',
       lambda: (mrg(OH2, [['C1', 100, '张三']], ['小班号', '验收人'], [['C1', '李四']])[1]['dropped_nonempty'],
                mrg(OH2, [['C1', 100, '张三']], ['小班号', '验收人'], [['C1', '李四']])[1]['old_only_cols']),
       lambda r: r == (1, ['小班面积']))

OHD = ['小班号', '验收人']
passes('F11 新表唯一键重复 → 计数并当新行追加（不猜更新哪一行）',
       lambda: (mrg(OHD, [['D1', '甲']], OHD, [['D1', ''], ['D1', '']])[1]['dup_new_key_count'],
                len(mrg(OHD, [['D1', '甲']], OHD, [['D1', ''], ['D1', '']])[0])),
       lambda r: r == (1, 2))
passes('F12 旧表本身有重复键 → 只合并一次，其余按"仅旧数据"保留',
       lambda: (mrg(OH, [['E1', 1, 'a', 'r1'], ['E1', 2, 'b', 'r2']], NH,
                    [['E1', 'x', 'y', 9, 'N']])[0][0],
                mrg(OH, [['E1', 1, 'a', 'r1'], ['E1', 2, 'b', 'r2']], NH,
                    [['E1', 'x', 'y', 9, 'N']])[1]['kept_old_only']),
       lambda r: r == (['E1', 'a', 'r1', 9, 'N'], 1))
raises('F13 唯一键列在任一侧缺失 → 报错（两边都得有）',
       lambda: mrg(['小班面积'], [['x']], NH, NROWS, key='小班号'), '都存在')

# ── 导出筛选候选值的三路合并（G，v0.26.1）───────────────────
print('\n=== G. 筛选下拉候选值：三路合并 ===')
GH = ['标段', '乡镇', '验收人', '验收结果']
GROWS = [['7标段', '羊街乡', '', ''],
         ['8标段', '羊街乡', '', '合格']]
GROWS2 = [r + [''] for r in GROWS]
GH2 = GH + ['检查员']
GCFG = {'验收结果选项': '合格;不合格'}
GUSERS = ['雷华雄', '苏正鹏', '何明星']

passes('G1 参数枚举并入（还没人填过也能选）',
       lambda: filter_options(GROWS, GH, '验收结果', GCFG, GUSERS),
       lambda r: r == ['合格', '不合格'])
passes('G2 含「人」的列并入系统用户名单',
       lambda: filter_options(GROWS, GH, '验收人', GCFG, GUSERS),
       lambda r: r == GUSERS)
passes('G3 含「员」的列同样并入用户名单',
       lambda: filter_options(GROWS2, GH2, '检查员', GCFG, GUSERS), lambda r: r == GUSERS)
passes('G4 不含「人/员」的列不并入用户（只取数据）',
       lambda: filter_options(GROWS, GH, '标段', GCFG, GUSERS), lambda r: r == ['7标段', '8标段'])
passes('G5 数据值与枚举重复时去重',
       lambda: filter_options(GROWS, GH, '验收结果', GCFG, GUSERS).count('合格'), lambda r: r == 1)
passes('G6 不传 cfg/users → 退化为只看数据（老行为）',
       lambda: filter_options(GROWS, GH, '验收结果'), lambda r: r == ['合格'])
passes('G7 无参数无用户时仍是数据首现顺序',
       lambda: filter_options(GROWS, GH, '乡镇', GCFG, GUSERS), lambda r: r == ['羊街乡'])

# ── H. 云同步远端路径规范化（v0.29.2）────────────────────────
# 起因：用户把「网盘应用目录」填成 /apps/x/ 时，面板显示与拼接会出现 //
print('\n=== H. 云同步远端路径：折叠重复斜杠 ===')
import sync_cloud as _SC  # noqa: E402


def _pp(app_dir, prefix, subdir='data/YJ-515', fname='a.jpg'):
    return _SC.photo_paths({'sync_baidu_app_dir': app_dir, 'sync_baidu_prefix': prefix},
                           subdir, fname)


passes('H1 应用目录尾部带斜杠 → 不产生 //',
       lambda: _pp('/apps/supervision/', 'supervision'),
       lambda r: r[1] == '/apps/supervision/supervision/data/YJ-515/a.jpg')
passes('H2 用户填了连续斜杠 → 折叠为单斜杠',
       lambda: _pp('/apps//x//', '/sub//d/'),
       lambda r: r == ('sub/d/data/YJ-515/a.jpg', '/apps/x/sub/d/data/YJ-515/a.jpg'))
passes('H3 七牛 key 不以 / 开头', lambda: _pp('/apps/x', '/p/')[0][0] != '/', lambda r: r is True)
passes('H4 无子目录时路径正确',
       lambda: _pp('/apps/book_translator', 'supervision', subdir='')[1],
       lambda r: r == '/apps/book_translator/supervision/a.jpg')
passes('H5 历史成功配置的结构保持一致',
       lambda: _pp('/apps/book_translator', 'supervision')[1],
       lambda r: r == '/apps/book_translator/supervision/data/YJ-515/a.jpg')
passes('H6 「网盘子目录」留空 → 不再多一层（去掉重复的 supervision）',
       lambda: _pp('/apps/supervision/', '')[1],
       lambda r: r == '/apps/supervision/data/YJ-515/a.jpg')

# ── I. 工作簿级保留占位符（v0.29.3）───────────────────────────
# 目的：让"按工作簿分层"由模板配置决定，而不是把目录结构写死在代码里
print('\n=== I. 占位符保留字：工作簿id / 工作簿名 / 工作簿文件名 ===')
HDRS = ['小班号', '标段']
passes('I1 目录模板可用 {{工作簿id}}_{{工作簿名}}/{{小班号}} → 校验通过',
       lambda: parse_params(prows(
           ['可编辑列', '小班号'], ['功能', '轨迹'],
           ['相片备注', '{{工作簿名}}'], ['相片文件名', '{{工作簿id}}_{{小班号}}'],
           ['目录', '{{工作簿id}}_{{工作簿名}}/{{小班号}}']), HDRS)['目录'],
       lambda r: r == '{{工作簿id}}_{{工作簿名}}/{{小班号}}')
passes('I2 {{工作簿文件名}} 也算保留字',
       lambda: parse_params(prows(
           ['可编辑列', '小班号'], ['功能', '轨迹'],
           ['相片备注', '{{工作簿文件名}}'], ['相片文件名', '{{小班号}}'],
           ['目录', '{{工作簿文件名}}/{{小班号}}']), HDRS)['目录'],
       lambda r: r == '{{工作簿文件名}}/{{小班号}}')
raises('I3 未知占位符仍然报错（保留字没有放宽校验）',
       lambda: parse_params(prows(
           ['可编辑列', '小班号'], ['功能', '轨迹'],
           ['相片备注', '{{小班号}}'], ['相片文件名', '{{小班号}}'],
           ['目录', '{{工作簿编号}}/{{小班号}}']), HDRS), '不是数据 sheet 表头')

# ── J. 验收流程解析（v0.31，doc/019）──────────────────────────────
# 重点覆盖两条实测踩过的坑：
#   J 组：括号感知分割（括号内的分号不能切）
#   K 组：自然语言附件解析（规则表顺序敏感、命中即停）
print('\n=== J. 流程解析：括号感知分割 + 控件结构 ===')
import flow_parser as fp  # noqa: E402

passes('J1 括号内分号不被切断：下拉三选项',
       lambda: fp.parse_ctl('完成情况=下拉(符合设计;部分完成;未完成)')[0],
       lambda r: r['name'] == '完成情况' and r['type'] == '下拉'
       and fp.ctl_options(r) == ['符合设计', '部分完成', '未完成'])
passes('J2 多选四选项完整（裸切会只剩第一个）',
       lambda: fp.ctl_options(fp.parse_ctl(
           '管护作业情况=多选(浇水;培土;补植;病虫害防治)')[0]),
       lambda r: r == ['浇水', '培土', '补植', '病虫害防治'])
passes('J3 括号外分号仍正常切分',
       lambda: [c['name'] for c in fp.parse_ctl(
           '底肥用量=数值(kg/袋,>0); 底肥品牌=单行文本(≤20字)')],
       lambda r: r == ['底肥用量', '底肥品牌'])
passes('J4 数值×3 产出三个子字段名（前后端共用，不各自拼字符串）',
       lambda: fp.parse_ctl('整地规格=数值×3(长/宽/深,cm,≥标准自动判)')[0]['sub'],
       lambda r: len(r) == 3 and all(x.startswith('整地规格') for x in r))
passes('J5 split_top 遇中文括号也不切',
       lambda: len(fp.split_top('下拉（甲;乙）; 数值(丙)', ';')),
       lambda r: r == 2)
passes('J6 「系统自动算」的字段带标记（后端据此免填）',
       lambda: '系统自动算' in fp.parse_ctl('成活率=数值(%,系统自动算)')[0]['arg'])
import openpyxl as _oxl  # noqa: E402


def _flow_sheet(file_name, sheet='流程-人工造林'):
    """按文件名找一份流程 Excel 并返回其 sheet 对象（read_only=True）。"""
    import pathlib
    for p in sorted(pathlib.Path(BASE).glob('data/**/*.xlsx')):
        if 'workbooks' in p.parts or p.name.startswith('.~'):
            continue
        try:
            wb = _oxl.load_workbook(p, read_only=True)
        except Exception:
            continue
        if sheet in wb.sheetnames and p.name == file_name:
            return wb, wb[sheet]
        wb.close()
    return None, None


_WB, _WS = _flow_sheet('造林项目-验收模板-含验收流程.xlsx')
if _WS is not None:
    # read_only 工作簿流**只能顺序读一次** —— 所以这里在同一次调用里把表头和数据一起拿，
    # 不能为了"再看一眼表头"再 iter_rows 一次（那会拿到空值 → 误报缺列，2026-10-14 踩过）
    passes('J8 read_only 下表头与数据一次读全（表头含「填报控件」）',
           lambda: '填报控件' in fp.parse_flow_sheet(_WS)[0]['controls_raw'] or
                   len(fp.parse_flow_sheet(_WS)[0]['controls']) > 0,
           lambda r: bool(r))
    _wb2 = _oxl.load_workbook(
        BASE / 'data/验收流程设计/造林项目-验收模板-含验收流程 副本.xlsx',
        read_only=True)
    if '流程-人工造林' in _wb2.sheetnames:
        passes('J9 真的缺「填报控件」列 → 每个环节带 warn',
               lambda: all(s.get('warn') for s in fp.parse_flow_sheet(_wb2['流程-人工造林'])),
               lambda r: bool(r))
    _wb2.close()

print('\n=== K. 流程解析：附件要求两层语法 ===')
passes('K1 正式语法 拍照:xx@2 → 下限 2 张',
       lambda: fp.parse_att('拍照:整地现场照片@2')[0],
       lambda r: r['kind'] == 'photo' and r['label'] == '整地现场照片' and r['min'] == 2)
passes('K2 正式语法 半角分号组合出两个按钮',
       lambda: [x['kind'] for x in fp.parse_att('上传:三证一签;拍照:剪苗作业')],
       lambda r: r == ['upload', 'photo'])
passes('K3 自然语言「各 1 张」拆成 3 个按钮且只有最后一个带下限',
       lambda: fp.parse_att('回塘、放肥、覆土各 1 张照片'),
       lambda r: [x['label'] for x in r] == ['回塘', '放肥', '覆土']
       and [x['min'] for x in r] == [0, 0, 1])
passes('K4 兜底规则不产生冗余按钮（每环节只一个"现场"）',
       lambda: [x['label'] for x in fp.parse_att('成活率抽查照片')],
       lambda r: r == ['成活率抽查'])
passes('K5 定位打卡 → punch 类型（不是 photo）',
       lambda: fp.parse_att('定位打卡 + 小班边界照片'),
       lambda r: r[0]['kind'] == 'punch' and r[0]['min'] == 1 and r[1]['kind'] == 'photo')
passes('K6 证件类 → upload（下拉/上传按钮类型正确）',
       lambda: [x['kind'] for x in
                fp.parse_att('三证一签（需良种证的树种传四证一签）、种苗规格照片、剪苗作业照片')],
       lambda r: r == ['upload', 'photo', 'photo'])
passes('K7 省略写法并入上一片（撑杆、覆膜照片 → 2 个按钮）',
       lambda: [x['label'] for x in fp.parse_att('撑杆、覆膜照片')],
       lambda r: r == ['撑杆', '覆膜'])
passes('K8 空附件要求 → 无按钮',
       lambda: fp.parse_att(''), lambda r: r == [])

print('\n=== L. 流程解析：标准值 trim 取值 + 流程匹配 ===')
passes('L1 「总株数  」尾随 2 空格，trim 后仍能取到（铁律 3）',
       lambda: fp.std_lookup({'总株数  ': 2960, '小班号': 'YJ-1'}).get('总株数'),
       lambda r: r == 2960)
passes('L2 std_lookup 同时保留原名与 trim 名',
       lambda: sorted(fp.std_lookup({'总株数  ': 1}).keys()),
       lambda r: r == ['总株数', '总株数  '])

REG = [{'sheet': '流程-人工造林', 'enabled': True, 'cond_field': '项目类型',
        'cond_val': '人工造乔木林', 'scope': '全部', 'prio': 10, 'line': 2},
       {'sheet': '流程-退化林修复', 'enabled': True, 'cond_field': '项目类型',
        'cond_val': '退化林修复', 'scope': '全部', 'prio': 20, 'line': 3},
       {'sheet': '流程-2026', 'enabled': False, 'cond_field': '项目类型',
        'cond_val': '人工造乔木林', 'scope': '全部', 'prio': 5, 'line': 5}]
FLOWS = {k: [{'no': 1, 'step': k, 'controls': [], 'attachments': []}]
         for k in ('流程-人工造林', '流程-退化林修复', '流程-2026')}
passes('L3 条件匹配：项目类型=人工造乔木林 → 命中人工造林',
       lambda: fp.match_flow(REG, FLOWS, {'项目类型': '人工造乔木林'})['sheet'],
       lambda r: r == '流程-人工造林')
passes('L4 值先 trim 再比（data 有尾随空格脏数据）',
       lambda: fp.match_flow(REG, FLOWS, {'项目类型': '人工造乔木林   '})['sheet'],
       lambda r: r == '流程-人工造林')
passes('L5 停用的流程不参与匹配（启用=否）',
       lambda: fp.match_flow(REG, FLOWS, {'项目类型': '人工造乔木林'})['sheet'],
       lambda r: r != '流程-2026')
passes('L6 条件值不命中 → None（前端退回网格填表）',
       lambda: fp.match_flow(REG, FLOWS, {'项目类型': '低效林改造'}),
       lambda r: r is None)
passes('L7 适用范围限定到具体文件时不生效',
       lambda: fp.match_flow(
           [{'sheet': '流程-人工造林', 'enabled': True, 'cond_field': '1',
             'cond_val': '', 'scope': '别的文件.xlsx', 'prio': 10, 'line': 2}],
           FLOWS, {'项目类型': 'x'}, '造林项目-验收模板.xlsx'),
       lambda r: r is None)
passes('L8 多条件全部满足才命中',
       lambda: fp.match_flow(
           [{'sheet': '流程-人工造林', 'enabled': True, 'cond_field': '项目类型;实施年度',
             'cond_val': '人工造乔木林;2026年', 'scope': '全部', 'prio': 10, 'line': 2}],
           FLOWS, {'项目类型': '人工造乔木林', '实施年度': '2027年'}),
       lambda r: r is None)
passes('L9 多条件都满足则命中',
       lambda: fp.match_flow(
           [{'sheet': '流程-人工造林', 'enabled': True, 'cond_field': '项目类型;实施年度',
             'cond_val': '人工造乔木林;2026年', 'scope': '全部', 'prio': 10, 'line': 2}],
           FLOWS, {'项目类型': '人工造乔木林', '实施年度': '2026年'})['sheet'],
       lambda r: r == '流程-人工造林')

print('\n=== M. 流程解析：多 sheet 数据表定位（v0.31）===')
passes('M1 8 个 sheet（含流程/说明）也能定位 data',
       lambda: _pick_data_sheet(['使用说明', 'data', '参数', '验收流程',
                                 '流程-人工造林', '流程-退化林修复',
                                 '流程-四旁植树', '因子控件对照']),
       lambda r: r == 'data')
raises('M2 只有流程 sheet 没有 data → 明确报错（不静默取流程表当数据）',
       lambda: _pick_data_sheet(['参数', '验收流程', '流程-人工造林']),
       '没有数据 sheet')
passes('M3 说明 sheet 不算数据表（只剩 data 一个）',
       lambda: _pick_data_sheet(['使用说明', 'data', '参数', '因子控件对照']),
       lambda r: r == 'data')
# 「混交比例」列 number_format 被误设 h:mm → openpyxl 读成 datetime.time，
# 不处理就会落库成 '07:03:00'，界面上"混交比例 07:03:00"（2026-10-14 浏览器实测发现）
passes('M4 被误设 h:mm 的数值列还原成小数（不是时间字符串）',
       lambda: _norm_cell(__import__('datetime').time(7, 3, 0)),
       lambda r: isinstance(r, float) and 0.29 < r < 0.30)
passes('M5 真日期仍走 date 分支（不被 time 分支抢走）',
       lambda: _norm_cell(__import__('datetime').date(2026, 10, 14)),
       lambda r: r == '2026-10-14')
passes('M6 datetime 带时分仍完整输出',
       lambda: _norm_cell(__import__('datetime').datetime(2026, 10, 14, 9, 30, 15)),
       lambda r: r == '2026-10-14 09:30:15')

print('\n' + '=' * 62)
print(f'通过 {_pass} / 失败 {_fail}')
sys.exit(1 if _fail else 0)
