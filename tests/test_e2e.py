#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""端到端测试（v0.23/v0.24，C11/C13）：真实模板解析 + 完整 HTTP 链路 + 下架语义。

用法：python tests/test_e2e.py      # 全绿退出码 0

设计原则：**不写死模板文件名与期望结果**。
- 模板从 data/*.xlsx 自动发现（用户会改名/增删），只认含「参数」sheet 的；
- "应当被拒/通过"从文件内容独立推断（见 file_expectation），模板被编辑不会让测试假失败；
- 依赖不满足时标 SKIP 而非失败。

注意：会在本地库 data/app.sqlite3 里临时建工作簿，跑完自动删除。
"""
import datetime
import io
import shutil
import sys
import tempfile
from pathlib import Path
from urllib.parse import urlencode

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

import openpyxl  # noqa: E402
import openpyxl.utils  # noqa: E402

from main import UPLOAD_DIR, ParamError, create_app, parse_workbook_storage  # noqa: E402

DATA = BASE / 'data'
ADMIN = ('雷华雄', 'lhx123')      # 文档记载的管理员
PLAIN = ('何明星', 'hmx123')      # 普通用户（若改过密码会 SKIP 相关断言）

_pass = _fail = _skip = 0


def ok(name, cond, detail=''):
    global _pass, _fail
    _pass, _fail = _pass + bool(cond), _fail + (not cond)
    print(f'  {"✅" if cond else "❌"} {name}  {detail}')


def skip(name, why):
    global _skip
    _skip += 1
    print(f'  ⏭  {name}  （SKIP：{why}）')


def _txt(v):
    return '' if v is None else str(v).strip()


def sheets_of(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def parse(path):
    """→ 结果 dict，或 ('ERR', 错误消息)。"""
    try:
        sn, h, rows, cfg, pr, kc = parse_workbook_storage(open(path, 'rb'))
        return {'sheet': sn, 'cols': len(h), 'rows': len(rows), 'key': kc, 'cfg': cfg}
    except ParamError as e:
        return ('ERR', str(e))


def file_expectation(path):
    """独立于被测代码，从 xlsx 内容推断"后端应当接受还是拒绝"。

    返回 'accept' / 'skip' / ('reject', 错误里应出现的关键词)。
    """
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        names = wb.sheetnames
        if '参数' not in names:
            return ('reject', '参数') if 'data' in names else 'skip'
        others = [n for n in names if n != '参数']
        if not others:
            return ('reject', '数据 sheet')
        if 'data' in others:
            dname = 'data'
        elif len(others) > 1:
            return ('reject', '无法判断')
        else:
            dname = others[0]

        head = [_txt(h) for h in (next(wb[dname].iter_rows(values_only=True), None) or ())]
        while head and head[-1] == '':
            head.pop()

        uk = None
        for row in wb['参数'].iter_rows(min_col=1, max_col=2, values_only=True):
            if _txt(row[0]).lower() == 'unique-key':
                uk = _txt(row[1])
        if not uk:
            return 'accept'                     # 未声明唯一键 → 不校验
        if uk not in head:
            return ('reject', '不在数据 sheet 表头')
        ci = head.index(uk)
        seen = set()
        for row in wb[dname].iter_rows(min_row=2, values_only=True):
            v = row[ci] if ci < len(row) else None
            if _txt(v) == '':
                return ('reject', '为空')
            k = _txt(v)
            if k in seen:
                return ('reject', '重复')
            seen.add(k)
        return 'accept'
    finally:
        wb.close()


def is_template(path):
    """含「参数」sheet，且数据 sheet 明确（名为 data，或只有一个非参数 sheet）。"""
    try:
        names = sheets_of(path)
    except Exception:
        return False
    if '参数' not in names:
        return False
    others = [n for n in names if n != '参数']
    return bool(others) and ('data' in others or len(others) == 1)


ALL_XLSX = sorted(DATA.glob('*.xlsx')) if DATA.is_dir() else []
TEMPLATES = [p for p in ALL_XLSX if is_template(p)]
DEFINED = [(p, sheets_of(p)) for p in TEMPLATES]
# 合成用例的基底：优先用有 data 表的模板（改名/加列更可控）
BASE_TPL = next((p for p, s in DEFINED if 'data' in s), TEMPLATES[0] if TEMPLATES else None)

print(f'=== 发现模板 {len(TEMPLATES)} 个 ===')
for p, s in DEFINED:
    print(f'    {p.name}  sheets={s}')


def data_sheet_of(wb):
    return 'data' if 'data' in wb.sheetnames else [n for n in wb.sheetnames if n != '参数'][0]


def set_unique_key(wb, col):
    """把参数 sheet 的 unique-key 指向 col；没有该行就补一行。"""
    wsp = wb['参数']
    for r_ in range(1, wsp.max_row + 1):
        if str(wsp.cell(r_, 1).value or '').strip().lower() == 'unique-key':
            wsp.cell(r_, 2).value = col
            return
    wsp.cell(wsp.max_row + 1, 1).value = 'unique-key'
    wsp.cell(wsp.max_row, 2).value = col


def drop_unique_key(wb):
    wsp = wb['参数']
    for r_ in range(wsp.max_row, 0, -1):
        if str(wsp.cell(r_, 1).value or '').strip().lower() == 'unique-key':
            wsp.delete_rows(r_)
            return


# ── A. 真实模板解析（期望由文件内容推断）────────────────────
print('\n=== A. 真实模板解析（期望由文件内容推断）===')
if not TEMPLATES:
    skip('A', 'data/ 下没有可用模板')
for p, _s in DEFINED:
    exp = file_expectation(p)
    r = parse(p)
    got_err = isinstance(r, tuple)
    if exp == 'skip':
        skip(f'A {p.name}', '不是填表模板')
    elif exp == 'accept':
        ok(f'A {p.name}：推断=应通过 → 实际通过', not got_err and bool(r['key']),
           r[1] if got_err else f"key={r['key']} 行={r['rows']}")
    else:
        ok(f'A {p.name}：推断=应被拒（含「{exp[1]}」）→ 实际被拒',
           got_err and exp[1] in r[1], r[1] if got_err else f'意外通过：{r}')

# ── B. 多 sheet 歧义（合成）─────────────────────────────────
print('\n=== B. 多 sheet 歧义 ===')
tmp = Path(tempfile.mkdtemp())
if BASE_TPL is None:
    skip('B', '无可用基底模板')
else:
    amb = tmp / 'amb.xlsx'
    shutil.copy(BASE_TPL, amb)
    wb = openpyxl.load_workbook(amb)
    for _ws in wb.worksheets:                 # 不用 wb.active：它不一定指向数据表
        if _ws.title != '参数':
            _ws.title = '甲'
            break
    wb.create_sheet('乙')
    wb.save(amb)
    wb.close()
    r = parse(amb)
    ok('B1 无 data 表且多个非参数 sheet → 明确报错并列出 sheet 名',
       isinstance(r, tuple) and '无法判断' in r[1] and '甲' in r[1],
       r[1] if isinstance(r, tuple) else r)

    okn = tmp / 'okname.xlsx'
    shutil.copy(BASE_TPL, okn)
    wb = openpyxl.load_workbook(okn)
    for _ws in wb.worksheets:
        if _ws.title != '参数':
            _ws.title = 'data'
            break
    wb.create_sheet('说明')
    drop_unique_key(wb)          # 摘掉唯一键，单独验证"多余 sheet 可容忍"
    wb.save(okn)
    wb.close()
    r = parse(okn)
    ok('B2 数据表名为 data + 多一个「说明」页 → 正常解析',
       not isinstance(r, tuple) and r['sheet'] == 'data',
       r[1] if isinstance(r, tuple) else f"sheet={r['sheet']} key={r['key']}")

# ── C. HTTP 完整链路 ────────────────────────────────────────
print('\n=== C. HTTP 完整链路（Flask test_client）===')
app = create_app()
# 两个身份必须用**两个独立 client**：Flask test_client 自带 cookie jar，
# 而 auth_user() 优先读 session、X-Sup-Token 只是后备 —— 同一 client 先后登录
# 会让后登录者的 session 覆盖前者，导致"管理员"请求被判成普通用户。
ca = app.test_client()      # 管理员
cu = app.test_client()      # 普通用户
CREATED = []                # 测试期间新建的工作簿 id，跑完统一 purge


def login(cli, user, pwd):
    r = cli.post('/api/login', json={'username': user, 'password': pwd})
    return r.status_code, (r.get_json() or {}).get('token', '')


_code, TA = login(ca, *ADMIN)
ok('C1 管理员登录', _code == 200 and bool(TA), f'HTTP {_code}')
_code_u, TU = login(cu, *PLAIN)
if not TU:
    skip('C1b 普通用户登录', f'{PLAIN[0]} 登录失败（可能改过密码）')
HA = {'X-Sup-Token': TA}
HU = {'X-Sup-Token': TU} if TU else None


def upload(path, name, headers=None, cli=None):
    with open(path, 'rb') as f:
        r = (cli or ca).post('/api/workbooks', data={'file': (f, name)},
                             headers=headers or HA, content_type='multipart/form-data')
    jj = r.get_json() or {}
    if jj.get('id'):
        CREATED.append(jj['id'])
    return r


def list_ids(cli, headers):
    body = cli.get('/api/workbooks', headers=headers).get_json() or {}
    return sorted(w['id'] for w in body.get('workbooks', []))


if BASE_TPL is None:
    skip('C2~C10', '无可用基底模板')
else:
    # C2 合成"可上传且唯一键指向新列"的模板：验证参数化真的生效
    good = tmp / 'good.xlsx'
    shutil.copy(BASE_TPL, good)
    wb = openpyxl.load_workbook(good)
    ws = wb[data_sheet_of(wb)]
    ci = ws.max_column + 1
    ws.cell(1, ci).value = '地块号'
    for r_ in range(2, ws.max_row + 1):
        ws.cell(r_, ci).value = f'DK-{r_ - 1:05d}'
    set_unique_key(wb, '地块号')
    wb.save(good)
    wb.close()

    rr = upload(good, 'good.xlsx')
    jj = rr.get_json() or {}
    ok('C2 上传（unique-key=地块号，数据唯一）→ 通过',
       rr.status_code == 200 and jj.get('ok'), f'HTTP {rr.status_code} {jj}')
    wid = jj.get('id')

    if wid:
        d = ca.get(f'/api/workbooks/{wid}', headers=HA).get_json() or {}
        ok('C3 返回 key_column = 地块号（参数化生效）', d.get('key_column') == '地块号',
           f"key_column={d.get('key_column')!r}")
        ok('C4 返回 log_fields（前端联动用）',
           isinstance(d.get('log_fields'), list) and bool(d['log_fields']),
           f"log_fields={d.get('log_fields')}")
        rr = ca.get(f'/api/workbooks/{wid}/export', headers=HA)
        ok('C5 导出仍是真 xlsx', rr.status_code == 200 and rr.data[:2] == b'PK',
           f'HTTP {rr.status_code} {len(rr.data)}B')
        rr = ca.post(f'/api/workbooks/{wid}/rows/0', json={'values': d['rows'][0]}, headers=HA)
        ok('C6 单行保存仍可用', rr.status_code == 200, f'HTTP {rr.status_code}')
        ok('C7 导出后再彻底删除 → 源模板目录被清',
           ca.delete(f'/api/workbooks/{wid}?purge=1', headers=HA).status_code == 200
           and not (UPLOAD_DIR / str(wid)).exists())

    # C8 唯一键重复 → 被拒（合成：新增一列，**先全部填满唯一值**再人为制造一对重复，
    #     否则会先被"为空"拦下，测不到重复逻辑）
    dup = tmp / 'dup.xlsx'
    shutil.copy(BASE_TPL, dup)
    wb = openpyxl.load_workbook(dup)
    ws = wb[data_sheet_of(wb)]
    ci = ws.max_column + 1
    ws.cell(1, ci).value = '测试键'
    for r_ in range(2, ws.max_row + 1):
        ws.cell(r_, ci).value = f'K-{r_ - 1:05d}'
    ws.cell(3, ci).value = ws.cell(2, ci).value          # 制造一对重复
    set_unique_key(wb, '测试键')
    wb.save(dup)
    wb.close()
    rr = upload(dup, 'dup.xlsx')
    jj = rr.get_json() or {}
    ok('C8 唯一键重复 → 被拒（带 Excel 行号）',
       rr.status_code == 400 and '重复' in jj.get('error', ''),
       f"HTTP {rr.status_code} {jj.get('error')}")

    # C9 未声明 unique-key → 不校验，通过且 key_column 回落
    nk = tmp / 'nokey.xlsx'
    shutil.copy(BASE_TPL, nk)
    wb = openpyxl.load_workbook(nk)
    drop_unique_key(wb)
    wb.save(nk)
    wb.close()
    rr = upload(nk, 'nokey.xlsx')
    jj = rr.get_json() or {}
    ok('C9 未声明 unique-key → 不校验，通过', rr.status_code == 200 and jj.get('ok'),
       f"HTTP {rr.status_code} {jj}")
    if jj.get('id'):
        d = ca.get(f"/api/workbooks/{jj['id']}", headers=HA).get_json() or {}
        ok('C9b 未声明时 key_column 回落小班号', d.get('key_column') == '小班号',
           f"key_column={d.get('key_column')!r}")
        ca.delete(f"/api/workbooks/{jj['id']}?purge=1", headers=HA)

    # C10 参数 sheet 被改名 → 被拒（合成）
    ren = tmp / 'renamed.xlsx'
    shutil.copy(BASE_TPL, ren)
    wb = openpyxl.load_workbook(ren)
    for _ws in wb.worksheets:
        if _ws.title == '参数':
            _ws.title = 'Sheet1'
            break
    wb.save(ren)
    wb.close()
    rr = upload(ren, 'renamed.xlsx')
    jj = rr.get_json() or {}
    ok('C10 「参数」sheet 被改名 → 被拒',
       rr.status_code == 400 and '参数' in jj.get('error', ''),
       f"HTTP {rr.status_code} {jj.get('error')}")

# ── D. 下架 / 上架 / 彻底删除（C13）─────────────────────────
print('\n=== D. 下架（软删除）/ 上架 / 彻底删除 ===')

if BASE_TPL is None:
    skip('D', '无可用基底模板')
else:
    base = tmp / 'inactive.xlsx'
    shutil.copy(BASE_TPL, base)
    wb = openpyxl.load_workbook(base)
    drop_unique_key(wb)
    wb.save(base)
    wb.close()

    rr = upload(base, 'inactive.xlsx')
    jj = rr.get_json() or {}
    wid = jj.get('id')
    ok('D1 上传一份干净模板', rr.status_code == 200 and bool(wid), f'HTTP {rr.status_code} {jj}')
    if wid:
        src = UPLOAD_DIR / str(wid)
        ok('D2 下架前：源模板目录存在', src.exists())
        if HU:
            ok('D3 下架前：普通用户在 App 列表里能看到它', wid in list_ids(cu, HU))

        rr = ca.delete(f'/api/workbooks/{wid}', headers=HA)
        ok('D4 管理员「下架」（DELETE 不带 purge）→ purged=false',
           rr.status_code == 200 and (rr.get_json() or {}).get('purged') is False,
           str(rr.get_json()))
        if HU:
            ok('D5 下架后：普通用户 App 列表里看不到它', wid not in list_ids(cu, HU), f'{list_ids(cu, HU)}')
            ok('D6 下架后：普通用户读详情 → 403',
               cu.get(f'/api/workbooks/{wid}', headers=HU).status_code == 403)
            ok('D7 下架后：普通用户单行保存 → 403',
               cu.post(f'/api/workbooks/{wid}/rows/0', json={'values': [1]},
                       headers=HU).status_code == 403)
            ok('D8 下架后：普通用户导出 → 403',
               cu.get(f'/api/workbooks/{wid}/export', headers=HU).status_code == 403)
        ok('D9 下架后：管理员仍可读详情（维护用）',
           ca.get(f'/api/workbooks/{wid}', headers=HA).status_code == 200)
        ok('D10 下架后：管理员仍可后台下载',
           ca.get(f'/admin/api/workbooks/{wid}/download', headers=HA).status_code == 200)
        ok('D11 下架后：源模板目录保留（可恢复）', src.exists())
        rows = (ca.get('/admin/api/workbooks', headers=HA).get_json() or {}).get('workbooks', [])
        me = [w for w in rows if w['id'] == wid]
        ok('D12 后台列表带 is_active=False', bool(me) and me[0]['is_active'] is False,
           str(me[:1])[:100])

        rr = ca.post(f'/api/workbooks/{wid}/restore', headers=HA)
        ok('D13 管理员「上架」→ 200', rr.status_code == 200)
        if HU:
            ok('D14 上架后：普通用户又能在列表看到它', wid in list_ids(cu, HU))
            ok('D15 上架后：普通用户可读详情',
               cu.get(f'/api/workbooks/{wid}', headers=HU).status_code == 200)
            ok('D16 普通用户无权下架 → 403',
               cu.delete(f'/api/workbooks/{wid}', headers=HU).status_code == 403)
            ok('D17 普通用户无权上架 → 403',
               cu.post(f'/api/workbooks/{wid}/restore', headers=HU).status_code == 403)

        rr = ca.delete(f'/api/workbooks/{wid}?purge=1', headers=HA)
        ok('D18 管理员「彻底删除」（purge=1）→ purged=true',
           rr.status_code == 200 and (rr.get_json() or {}).get('purged') is True,
           str(rr.get_json()))
        ok('D19 彻底删除后：源模板目录被清', not src.exists())
        ok('D20 彻底删除后：详情 404',
           ca.get(f'/api/workbooks/{wid}', headers=HA).status_code == 404)

# ── E. 导出筛选（后台「下载 Excel」弹框，只筛行）────────────
print('\n=== E. 导出筛选：声明 ∩ 表头 / 只筛行 / 日期真值 ===')
if BASE_TPL is None:
    skip('E', '无可用基底模板')
else:
    # 合成模板：摘掉 unique-key、加「导出筛选」声明、写入可控的前 3 行
    fsrc = tmp / 'filter.xlsx'
    shutil.copy(BASE_TPL, fsrc)
    wb = openpyxl.load_workbook(fsrc)
    ws = wb[data_sheet_of(wb)]
    heads = [_txt(c.value) for c in next(ws.iter_rows(max_row=1))]
    while heads and heads[-1] == '':
        heads.pop()

    def col(name):
        return heads.index(name) + 1 if name in heads else None

    cand = [f for f in ('标段', '乡镇', '验收人', '验收日期') if col(f)]
    if len(cand) < 4:
        skip('E', f'基底模板缺少字段（只有 {cand}）')
    else:
        drop_unique_key(wb)
        wsp = wb['参数']
        wsp.cell(wsp.max_row + 1, 1).value = '导出筛选'
        wsp.cell(wsp.max_row, 2).value = '标段|select;乡镇|select;验收人|select;验收日期|date'
        fake = [('甲标', 'A乡', '张三', '2026-09-10'),
                ('乙标', 'B乡', '李四', '2026-09-11'),
                ('甲标', 'A乡', '张三', '2026-09-12')]
        for k, (bd, xz, yr, dt) in enumerate(fake, start=2):
            ws.cell(k, col('标段')).value = bd
            ws.cell(k, col('乡镇')).value = xz
            ws.cell(k, col('验收人')).value = yr
            ws.cell(k, col('验收日期')).value = dt
        wb.save(fsrc)
        wb.close()

        rr = upload(fsrc, 'filter.xlsx')
        jj = rr.get_json() or {}
        wid = jj.get('id')
        ok('E1 上传带「导出筛选」的模板', rr.status_code == 200 and bool(wid), f'HTTP {rr.status_code} {jj}')

        if wid:
            fo = (ca.get(f'/admin/api/workbooks/{wid}/filter-options', headers=HA).get_json()
                  or {}).get('fields', [])
            ok('E2 filter-options = 声明的 4 个字段（顺序保持）',
               [x['field'] for x in fo] == ['标段', '乡镇', '验收人', '验收日期'],
               str([x['field'] for x in fo]))
            byf = {x['field']: x for x in fo}
            ok('E3 下拉字段带候选值且已去重',
               '甲标' in byf['标段']['options'] and '乙标' in byf['标段']['options'],
               str(byf['标段']['options'])[:90])
            ok('E4 日期字段 type=date 且无候选值',
               byf['验收日期']['type'] == 'date' and byf['验收日期']['options'] == [])

            def dl(qs=''):
                r = ca.get(f'/admin/api/workbooks/{wid}/download{qs}', headers=HA)
                assert r.status_code == 200, f'下载失败 HTTP {r.status_code}'
                return openpyxl.load_workbook(io.BytesIO(r.data))

            wbk = dl()
            ex = wbk[wbk.sheetnames[0]]
            ex_heads = [_txt(c.value) for c in next(ex.iter_rows(max_row=1))]
            last = openpyxl.utils.get_column_letter(len(heads))
            ok('E5 auto_filter 范围按实际数据重设',
               ex.auto_filter.ref == f'A1:{last}{ex.max_row}', repr(ex.auto_filter.ref))
            dcell = ex.cell(2, ex_heads.index('验收日期') + 1)
            ok('E6 验收日期导出为真日期单元格',
               isinstance(dcell.value, datetime.datetime), type(dcell.value).__name__)
            ok('E7 日期单元格带 yyyy-mm-dd 格式', dcell.number_format == 'yyyy-mm-dd',
               dcell.number_format)

            def first_rows(wbk_):
                ws_ = wbk_[wbk_.sheetnames[0]]
                hs_ = [_txt(c.value) for c in next(ws_.iter_rows(max_row=1))]
                out = []
                for r_ in range(2, ws_.max_row + 1):
                    if all(ws_.cell(r_, c_).value in (None, '') for c_ in range(1, ws_.max_column + 1)):
                        continue
                    out.append({h: ws_.cell(r_, hs_.index(h) + 1).value for h in cand})
                    if len(out) >= 3:
                        break
                return out

            full = dl()
            all_rows = first_rows(full)
            ok('E8 不带筛选参数 → 前 3 行原样导出', len(all_rows) == 3, str(all_rows)[:100])

            one = first_rows(dl(f'?{urlencode({"标段": "甲标"})}'))
            ok('E9 按 标段=甲标 筛 → 首行是甲标', one and one[0]['标段'] == '甲标', str(one[:2])[:100])

            und = first_rows(dl(f'?{urlencode({"小班号": "不存在"})}'))
            ok('E10 未声明的字段被忽略（不会把数据筛空）', len(und) >= 1, f'{len(und)} 行')

            dat = first_rows(dl(f'?{urlencode({"验收日期": "2026-09-11"})}'))
            ok('E11 按日期单日筛 → 命中 2026-09-11 那一行',
               len(dat) == 1 and str(dat[0]['验收日期'])[:10] == '2026-09-11', str(dat))

            # 回传：导出的真日期重新上传应还原为 'YYYY-MM-DD' 文本
            buf = io.BytesIO()
            dl().save(buf)
            buf.seek(0)
            sn_, h_, rows_, cfg_, pr_, kc_ = parse_workbook_storage(buf)
            ci_ = h_.index('验收日期')
            ok('E12 导出文件重新上传 → 日期还原为 YYYY-MM-DD 文本',
               [rows_[k][ci_] for k in range(3)] == ['2026-09-10', '2026-09-11', '2026-09-12'],
               str([rows_[k][ci_] for k in range(3)]))

            # 未声明筛选的模板 → 空清单（前端就不弹框）
            nf = tmp / 'nofilter.xlsx'
            shutil.copy(BASE_TPL, nf)
            wb = openpyxl.load_workbook(nf)
            drop_unique_key(wb)
            wb.save(nf)
            wb.close()
            rr = upload(nf, 'nofilter.xlsx')
            jj2 = rr.get_json() or {}
            if jj2.get('id'):
                fo2 = (ca.get(f"/admin/api/workbooks/{jj2['id']}/filter-options",
                              headers=HA).get_json() or {}).get('fields')
                ok('E13 未声明「导出筛选」→ 空清单', fo2 == [], str(fo2))

            # /admin 跳转（v0.25）：未登录带回 next=admin；非管理员回 App
            cn = app.test_client()          # 全新的未登录 client
            r1 = cn.get('/admin')
            ok('E14 未登录访问 /admin → 302 且带 next=admin',
               r1.status_code == 302 and 'next=admin' in (r1.headers.get('Location') or ''),
               f"{r1.status_code} {r1.headers.get('Location')}")
            if HU:
                r2 = cu.get('/admin')
                ok('E15 非管理员访问 /admin → 302 且不带 next',
                   r2.status_code == 302 and 'next=' not in (r2.headers.get('Location') or ''),
                   f"{r2.status_code} {r2.headers.get('Location')}")
            r3 = ca.get('/admin')
            ok('E16 管理员访问 /admin → 200', r3.status_code == 200, f'HTTP {r3.status_code}')

# ── 清理本次测试新建的工作簿（彻底删除，含源模板目录）──────
print('\n=== 清理测试数据 ===')
for _wid in sorted(set(CREATED)):
    _r = ca.delete(f'/api/workbooks/{_wid}?purge=1', headers=HA)
    print(f'    purge id={_wid} → HTTP {_r.status_code}')

print('\n' + '=' * 62)
print(f'通过 {_pass} / 失败 {_fail} / 跳过 {_skip}')
sys.exit(1 if _fail else 0)
