"""不变量检查：把业务勾稽关系写成断言，每次生成输出后必跑。
用法: python check_invariants.py 输出目录/

与 compare.py 分工：
- compare.py 证明「没改坏旧行为」（新旧双跑 diff，重构时用）
- 本脚本证明「输出满足业务勾稽」（如 合计=明细和、A列=B列-C列，改口径时两者都要）

按项目填 CHECKS：每项 = (名称, 函数(ctx) -> (ok: bool, 详情: str))。
退出码非 0 = 有违反 → 可挂 pre-commit/CI。"""
import sys
from pathlib import Path

TOL = 1e-4  # 数值容差，与 compare.py 一致

class Ctx:
    """惰性加载输出文件 + 常用断言助手。按输出格式扩展。"""
    def __init__(self, out_dir: Path):
        self.dir = out_dir
        self._wb = {}

    def cell(self, file, sheet, coord):
        """读 Excel 单元格值。"""
        from openpyxl import load_workbook
        if file not in self._wb:
            self._wb[file] = load_workbook(self.dir / file, data_only=True)
        return self._wb[file][sheet][coord].value

    def near(self, a, b, tol=TOL):
        """数值近似相等。"""
        try:
            return abs(float(a) - float(b)) <= tol
        except (TypeError, ValueError):
            return False

# ── 按项目填：业务勾稽断言 ─────────────────────────
def _示例合计等于明细之和(c: Ctx):
    """示例：E6 = E7+E8+E9+E10。换成你的文件/sheet/单元格后取消 CHECKS 注释。"""
    f, s = "主表.xlsx", "Sheet1"
    total = c.cell(f, s, "E6")
    parts = sum(c.cell(f, s, x) or 0 for x in ("E7", "E8", "E9", "E10"))
    return c.near(total, parts), f"E6={total} vs 明细和={parts}"

CHECKS = [
    # ("E6=明细和（示例，换成你的勾稽）", _示例合计等于明细之和),
]
# ────────────────────────────────────────────────────

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    if not CHECKS:
        print("CHECKS 为空：先把业务勾稽填进 CHECKS（文件内有示例）")
        sys.exit(2)
    ctx = Ctx(Path(sys.argv[1]))
    fails = 0
    for name, fn in CHECKS:
        try:
            ok, detail = fn(ctx)
        except Exception as e:  # 断言本身执行失败也计为违反，带出上下文
            ok, detail = False, f"断言执行异常: {e}"
        print(f"[{'PASS' if ok else 'FAIL'}] {name} | {detail}")
        fails += 0 if ok else 1
    print(f"\n不变量: {len(CHECKS)} 项 | 违反: {fails}")
    sys.exit(1 if fails else 0)
