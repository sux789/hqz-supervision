#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""模板体检：上传前先跑这个，把「参数」与「data」两个 sheet 的问题一次列全。

用法:
    python tools/check_template.py data/*.xlsx
    python tools/check_template.py data/xxx.xlsx --key 小班号   # 覆盖唯一键列名

三级结论:
    [FATAL]   上传会被拒 / 破坏「唯一键」假设 —— 必须修
    [PENDING] 模板写法正确，但后端还未实现该 key —— 属待开发的代码改动，不是模板的错
    [WARN]    不影响上传，建议修（数据脏值、写法不一致、体验问题）

退出码 0 = 无 FATAL，可安全上传。只读，不修改任何文件。
"""
import sys
from pathlib import Path
from collections import OrderedDict

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))
from param_parser import KEYS as BACKEND_KEYS  # noqa: E402
from param_parser import ParamError, parse_params  # noqa: E402

DEFAULT_KEY = '小班号'
DATA_SHEET = 'data'          # 约定的数据 sheet 名
PARAM_SHEET = '参数'
KEY_NAME = 'unique-key'      # 唯一键 key 名（后端只认这一个写法）
BACKEND_HAS_KEY = KEY_NAME in BACKEND_KEYS   # 后端是否已支持（未支持时报 PENDING 而非 FATAL）

NICE_KEYS = ['验收结果选项', '不显示列', '搜索选项', '压缩最长边', '压缩质量']
HEADER_COLS = 6
FULLWIDTH_WS = '\u3000'
SUSPECT_FW = {'｜': '|', '；': ';'}


def _s(v):
    return '' if v is None else str(v)


def check_file(path: Path, override_key=None):
    F, W, P, info = [], [], [], {}

    if path.suffix not in ('.xlsx', '.xlsm'):
        return ['只支持 .xlsx / .xlsm'], W, P, info

    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    names = wb.sheetnames
    info['sheets'] = names

    if PARAM_SHEET not in names:
        wb.close()
        # 区分「填表模板缺参数表」和「这压根不是填表模板」：
        # 有名为 data 的 sheet → 是模板（漏了改名/漏了参数表）；否则视为非模板文件，跳过。
        if DATA_SHEET in names:
            return ([f'缺少「{PARAM_SHEET}」sheet（sheet 名必须精确为「{PARAM_SHEET}」，'
                     f'当前实为：{names}）'], W, P, info)
        info['skip'] = True
        return (F, W + [f'不是填表模板（无「{PARAM_SHEET}」也无「{DATA_SHEET}」sheet，'
                        f'当前为 {names}）→ 已跳过'], P, info)

    others = [n for n in names if n != PARAM_SHEET]
    if not others:
        wb.close()
        return ['只有「参数」sheet，没有数据 sheet'], W, P, info
    if len(others) > 1:
        F.append(f'有 {len(others)} 个非「参数」sheet：{others}。后端用「第一个非参数 sheet」'
                 f'定位数据表（main.py:354），多 sheet 会取错 → 只保留 1 个数据 sheet')
    if DATA_SHEET in others:
        data_name = DATA_SHEET
    else:
        data_name = others[0]
        W.append(f'数据 sheet 名为「{data_name}」，非约定名「{DATA_SHEET}」'
                 f'（后端兼容，但通用化建议统一为「{DATA_SHEET}」）')
    info['data_sheet'] = data_name

    ws_p = wb[PARAM_SHEET]
    prows = [(i, r) for i, r in enumerate(
        ws_p.iter_rows(max_col=HEADER_COLS, values_only=True), 1)
        if any(x not in (None, '') for x in r)]

    ws_d = wb[data_name]
    raw_head = list(next(ws_d.iter_rows(values_only=True), ()) or ())
    while raw_head and raw_head[-1] in (None, ''):
        raw_head.pop()
    headers = [_s(h).strip() for h in raw_head]

    # ── 唯一键声明 ──────────────────────────────────────
    # 后端已支持 unique-key（v0.23+）→ 该行留在 kept 里交给 parse_params 正规校验；
    # 后端尚未支持（老版本代码配新模板）→ 摘出来报 PENDING，避免误判成模板的错。
    decl_key, kept = None, []
    for rno, cells in prows:
        c = [_s(x) for x in (list(cells) + [''] * HEADER_COLS)[:HEADER_COLS]]
        if c[0].strip().lower() == KEY_NAME:
            decl_key = c[1].strip()
            if BACKEND_HAS_KEY:
                kept.append((rno, cells))
            else:
                P.append(f'参数 sheet 第{rno}行声明「{KEY_NAME}」= {decl_key!r}，'
                         f'但后端 param_parser.KEYS 尚未收录该 key —— 属待开发，不是模板的错')
        else:
            kept.append((rno, cells))
    key_col = override_key or decl_key or DEFAULT_KEY
    info['key_col'] = key_col
    if not decl_key and not override_key:
        W.append(f'参数 sheet 未声明唯一键，回落到默认列「{DEFAULT_KEY}」')

    config = {}
    try:
        config = parse_params(kept, headers)
        info['config'] = config
    except ParamError as e:
        F.append(str(e))

    for k in NICE_KEYS:
        if k not in config:
            W.append(f'参数 sheet 未配置可选 key「{k}」'
                     + ('（照片压缩将回落后台默认 1440 / 0.8）' if k.startswith('压缩') else ''))

    for rno, cells in kept[1:]:
        c = [_s(x) for x in (list(cells) + [''] * HEADER_COLS)[:HEADER_COLS]]
        k, v = c[0].strip(), c[1]
        if not k or k == 'key':
            continue
        for ch, rep in SUSPECT_FW.items():
            if ch in v:
                F.append(f'参数 sheet 第{rno}行「{k}」value 含全角「{ch}」，应为半角「{rep}」')
        if '\\' in v:
            W.append(f'参数 sheet 第{rno}行「{k}」value 含反斜杠，疑似转义残留：{v!r}')

    # ── data 表表头 ─────────────────────────────────────
    if not headers:
        F.append(f'数据 sheet「{data_name}」缺少表头行')
        wb.close()
        return F, W, P, info

    # 表头首尾空格：后端 `parse_workbook_storage` 会用 _s() 先 strip 再匹配，
    # 所以**不影响占位符/可编辑列/搜索选项的匹配**——只报 WARN（Excel 里看不出来，易改错列）。
    # 只有当 strip 之后出现重名列时才是 FATAL。
    for i, h in enumerate(raw_head):
        if h not in (None, '') and _s(h) != _s(h).strip():
            W.append(f'data 表第 {i + 1} 列表头 {_s(h)!r} 含首尾空格'
                     f'（后端会 strip 后使用，不影响匹配；但 Excel 里看不出来，建议清掉）')
    blanks = [i + 1 for i, h in enumerate(raw_head) if h in (None, '') or _s(h).strip() == '']
    if blanks:
        W.append(f'data 表存在空表头列：{blanks}（建议删掉整列，避免列数错位）')
    dup_h = [h for h in set(headers) if h and headers.count(h) > 1]
    if dup_h:
        F.append(f'data 表列名重复（去空格后）：{dup_h} —— 取值与匹配会错乱，必须改名')

    ncols = len(headers)
    rows = []
    for r in ws_d.iter_rows(min_row=2, values_only=True):
        vals = list(r) + [None] * (ncols - len(r))
        if any(v not in (None, '') for v in vals[:ncols]):
            rows.append(vals[:ncols])
    info['rows'] = len(rows)

    # ── 唯一键校验 ──────────────────────────────────────
    if key_col not in headers:
        F.append(f'唯一键列「{key_col}」不在 data 表头中（表头：{"、".join(headers)}）')
    else:
        ki = headers.index(key_col)
        seen, dups, empties = {}, [], []
        for n, r in enumerate(rows, start=2):
            v = r[ki]
            if v in (None, ''):
                empties.append(n)
                continue
            k = _s(v).strip()
            if k in seen:
                dups.append((k, seen[k], n))
            else:
                seen[k] = n
        info['unique'] = (len(seen), len(dups), len(empties))
        if empties:
            F.append(f'唯一键「{key_col}」有 {len(empties)} 行为空：行 {empties[:6]}'
                     + (' …' if len(empties) > 6 else ''))
        for k, a, b in dups[:8]:
            F.append(f'唯一键「{key_col}」重复：{k!r} 出现在 行{a} 与 行{b}')
        if len(dups) > 8:
            F.append(f'唯一键「{key_col}」另有 {len(dups) - 8} 组重复未列出')
        editable = [x.strip() for x in config.get('可编辑列', '').split(';') if x.strip()]
        if key_col in editable:
            W.append(f'唯一键列「{key_col}」出现在「可编辑列」中 —— 一旦被改，行标识会漂移，'
                     f'建议从可编辑列里排除（C11）')

    # ── 脏值 ────────────────────────────────────────────
    for ci, h in enumerate(headers):
        raw_vals = [(_s(r[ci]), n) for n, r in enumerate(rows, start=2) if r[ci] not in (None, '')]
        sp = [(n, v) for v, n in raw_vals if v != v.strip()]
        fw = [(n, v) for v, n in raw_vals if FULLWIDTH_WS in v]
        groups = OrderedDict()
        for v, n in raw_vals:
            groups.setdefault(v.strip().replace(FULLWIDTH_WS, ''), OrderedDict()).setdefault(v, n)
        mixed = {k: d for k, d in groups.items() if len(d) > 1}
        if sp:
            W.append(f'列「{h}」有 {len(sp)} 个值带首尾空格，如 行{sp[0][0]}={sp[0][1]!r}')
        if fw:
            W.append(f'列「{h}」有 {len(fw)} 个值含全角空格，如 行{fw[0][0]}={fw[0][1]!r}')
        if mixed:
            detail = '；'.join(' / '.join(repr(x) for x in d)
                              for _, d in list(mixed.items())[:3])
            W.append(f'列「{h}」疑似同值异写 {len(mixed)} 组：{detail}')

    # ── 目录模板与固定 sheet 名的耦合 ───────────────────
    dircfg = config.get('目录', '')
    if '{{sheet名称}}' in dircfg and data_name == DATA_SHEET:
        W.append('「目录」用了 {{sheet名称}}，而数据 sheet 固定叫「data」'
                 '→ 照片子目录会变成 data/{唯一键}，建议改用 {{项目类型}} 或固定前缀')

    dvs = list(ws_d.data_validations.dataValidation) if ws_d.data_validations else []
    opts = [x.strip() for x in config.get('验收结果选项', '').split(';') if x.strip()]
    if opts and '验收结果' in headers:
        tgt = headers.index('验收结果') + 1
        covered = any(any(getattr(c, 'min_col', c[1]) <= tgt <= getattr(c, 'max_col', c[1])
                          for c in dv.sqref.ranges) for dv in dvs)
        if not covered:
            W.append(f'data 表「验收结果」列未绑定 Excel 下拉（参数已声明 {opts}；App 内有下拉，'
                     f'但直接在 Excel 里填会填出自由文本）')
    if not ws_d.freeze_panes:
        W.append('data 表未冻结首行（几百行滚动时容易看错列）')

    wb.close()
    return F, W, P, info


def main(argv):
    override_key = None
    if '--key' in argv:
        i = argv.index('--key')
        override_key = argv[i + 1]
        argv = argv[:i] + argv[i + 2:]
    files = [Path(a) for a in argv[1:]]
    if not files:
        print(__doc__)
        return 2

    total_f = 0
    for p in files:
        print('=' * 76)
        print(f'▶ {p.name}')
        try:
            F, W, P, info = check_file(p, override_key)
        except Exception as e:  # noqa: BLE001
            print(f'  [FATAL  ] 读取异常 {type(e).__name__}: {e}')
            total_f += 1
            continue
        if info.get('rows') is not None:
            u = info.get('unique')
            print(f'  sheets={info.get("sheets")}  data={info.get("data_sheet")!r}  '
                  f'唯一键={info.get("key_col")!r}  数据行={info.get("rows")}'
                  + (f'（唯一 {u[0]}）' if u else ''))
        skip = info.get('skip')
        for m in F:
            print(f'  [FATAL  ] {m}')
        for m in P:
            print(f'  [PENDING] {m}')
        for m in W:
            print(f'  {"[SKIP   ]" if skip else "[WARN   ]"} {m}')
        if skip:
            print('  —— 已跳过（不是填表模板）')
        elif not F and not W and not P:
            print('  ✅ 全部通过')
        else:
            print(f'  —— {len(F)} FATAL / {len(W)} WARN / {len(P)} PENDING')
        total_f += len(F)

    print('=' * 76)
    print(f'合计 {total_f} 个 FATAL —— ' + ('模板可安全上传' if total_f == 0 else '修完再上传'))
    return 1 if total_f else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
