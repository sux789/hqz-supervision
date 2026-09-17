# -*- coding: utf-8 -*-
"""hqz-supervision —— 通用 Excel 网格填表工具（B1 本地预览版）。

约束：C01 参数驱动 / C02 参数语法 / C03 不影响其他业务（本地 127.0.0.1:8720）
     C04 通用无调查业务 / C05 账号（雷华雄 admin）/ C06 水印日期不参数化
     C07 相片不落服务器（App 存系统相册 / 浏览器下载，前端仅记录文件名提示）
启动：python main.py  →  http://127.0.0.1:8720
"""
import base64
import hashlib
import hmac
import io
import json
import re
import shutil
import sqlite3
import threading
import zipfile
from datetime import date, datetime, time, timedelta
from functools import wraps
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from flask import (Blueprint, Flask, abort, jsonify, redirect, render_template,
                   request, send_file, session, url_for)

import sync_cloud
import track_export
from param_parser import (ParamError, _s, check_unique_column, export_filters_of,
                          filter_options, filter_rows, log_fields_of, merge_rows,
                          parse_params, percent_cols_of, row_key_column, split_list,
                          unique_key_of)

BASE = Path(__file__).resolve().parent
DATA_DIR = BASE / 'data'
UPLOAD_DIR = DATA_DIR / 'workbooks'
TRACK_DIR = DATA_DIR / 'tracks'
PENDING_DIR = DATA_DIR / 'pending'  # 云同步待推送缓冲（百度成功即删，C07 不长期存储）
DB_PATH = DATA_DIR / 'app.sqlite3'
for d in (UPLOAD_DIR, TRACK_DIR, PENDING_DIR):
    d.mkdir(parents=True, exist_ok=True)

PORT = 8720  # 独立端口（C03：不与 hqz-survey / hqz-cam-app 冲突）

bp = Blueprint('sup', __name__)

# ────────────────────────── 基础设施 ──────────────────────────

def db():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con


def init_db():
    con = db()
    con.executescript('''
        CREATE TABLE IF NOT EXISTS users(
            username TEXT PRIMARY KEY, password_hash TEXT NOT NULL, role TEXT NOT NULL,
            pwd_ver INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE IF NOT EXISTS workbooks(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL, sheet_name TEXT NOT NULL,
            headers TEXT NOT NULL, rows TEXT NOT NULL,
            config TEXT NOT NULL, param_rows TEXT NOT NULL,
            key_column TEXT NOT NULL DEFAULT '',
            is_active INTEGER NOT NULL DEFAULT 1,
            uploaded_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS accept_logs(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            workbook_id INTEGER NOT NULL,
            row_idx INTEGER NOT NULL,
            xiaoban TEXT,
            field TEXT NOT NULL,
            old_value TEXT, new_value TEXT,
            operator TEXT NOT NULL,
            created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS settings(
            key TEXT PRIMARY KEY, value TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS photo_sync(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            workbook_id INTEGER NOT NULL,
            xiaoban TEXT DEFAULT '',
            filename TEXT NOT NULL,
            remote_path TEXT NOT NULL,
            qiniu_key TEXT NOT NULL,
            size INTEGER NOT NULL DEFAULT 0,
            state TEXT NOT NULL DEFAULT 'received',
            retry_count INTEGER NOT NULL DEFAULT 0,
            last_error TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            synced_at TEXT, purged_at TEXT,
            kind TEXT NOT NULL DEFAULT 'photo',
            orig_size INTEGER NOT NULL DEFAULT 0);
        CREATE INDEX IF NOT EXISTS idx_photo_state ON photo_sync(state);
        -- v0.26：「更新 Excel」前的自动备份（DB 行 + 源模板文件），支持一键回滚
        CREATE TABLE IF NOT EXISTS workbook_backups(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            workbook_id INTEGER NOT NULL,
            name TEXT NOT NULL, sheet_name TEXT NOT NULL,
            headers TEXT NOT NULL, rows TEXT NOT NULL,
            config TEXT NOT NULL, param_rows TEXT NOT NULL,
            key_column TEXT NOT NULL DEFAULT '',
            src_ext TEXT NOT NULL DEFAULT '.xlsx',
            src_bak TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_wb_backup ON workbook_backups(workbook_id);
    ''')
    # 轻量迁移：v0.21 前建的 users 表无 pwd_ver 列（改密码令旧令牌失效用）
    ucols = {r[1] for r in con.execute('PRAGMA table_info(users)')}
    if 'pwd_ver' not in ucols:
        con.execute('ALTER TABLE users ADD COLUMN pwd_ver INTEGER NOT NULL DEFAULT 1')
    # 轻量迁移：v0.23 前建的 workbooks 无 key_column 列（唯一键列名，C11）
    wcols = {r[1] for r in con.execute('PRAGMA table_info(workbooks)')}
    if 'key_column' not in wcols:
        con.execute("ALTER TABLE workbooks ADD COLUMN key_column TEXT NOT NULL DEFAULT ''")
    # 轻量迁移：v0.24 前建的 workbooks 无 is_active 列（下架/上架，C13）
    if 'is_active' not in wcols:
        con.execute('ALTER TABLE workbooks ADD COLUMN is_active INTEGER NOT NULL DEFAULT 1')
    # 轻量迁移：v0.10 前建的 photo_sync 无 kind/orig_size 列
    cols = {r[1] for r in con.execute('PRAGMA table_info(photo_sync)')}
    if 'kind' not in cols:
        con.execute("ALTER TABLE photo_sync ADD COLUMN kind TEXT NOT NULL DEFAULT 'photo'")
    if 'orig_size' not in cols:
        con.execute('ALTER TABLE photo_sync ADD COLUMN orig_size INTEGER NOT NULL DEFAULT 0')
    if not con.execute('SELECT 1 FROM users').fetchone():
        con.execute(
            'INSERT INTO users VALUES(?,?,?)',
            ('雷华雄', hashlib.sha256('lhx123'.encode()).hexdigest(), 'admin', 1))
    # settings 非密钥默认值预置（真实 AK/SK / 百度凭证由一次性脚本写入，不进代码）
    for k, v in {**sync_cloud.SYNC_DEFAULTS, **sync_cloud.VIDEO_DEFAULTS, **sync_cloud.PHOTO_DEFAULTS}.items():
        con.execute('INSERT OR IGNORE INTO settings(key, value) VALUES(?,?)', (k, v))
    con.commit()
    con.close()


# ── 长期令牌（v0.15）──────────────────────────────────────────────
# 背景：系统相机是独立 Activity，会把 WebView 顶到后台甚至让进程被杀；此时
# WebView 的 cookie 可能尚未落盘 → 回来页面重载变"未登录"。令牌存 localStorage
# （写入即持久）+ 随每个请求带 X-Sup-Token，彻底摆脱 cookie 落盘时序问题。

_TOKEN_DAYS = 3650      # 10 年：App 登录后长期保存（改密码会令旧令牌立即失效）


def _token_secret() -> bytes:
    return b'hqz-supervision-token-v1'


def make_token(user: str, ver: int = 1, days: int = _TOKEN_DAYS) -> str:
    """令牌 = base64(user|密码版本|过期时间|签名)。

    带"密码版本"是为了让**改密码后旧令牌立即失效**（v0.21）：令牌无状态，无法逐条吊销，
    改密码时把 users.pwd_ver +1，旧令牌校验时版本对不上即拒绝。
    """
    exp = int(datetime.now().timestamp()) + days * 86400
    payload = f'{user}|{int(ver)}|{exp}'
    sig = hmac.new(_token_secret(), payload.encode(), hashlib.sha256).hexdigest()[:32]
    return b64e(f'{payload}|{sig}'.encode())


def verify_token(tok: str):
    """→ (用户名, 密码版本) 或 None（签名不符/过期/格式错）。"""
    try:
        parts = b64d(tok).decode().rsplit('|')
        if len(parts) == 3:            # 旧格式（v0.21 前，无密码版本）：按当前版本直接接受，避免升级即被登出
            user, exp, sig = parts
            if not hmac.compare_digest(
                    hmac.new(_token_secret(), f'{user}|{exp}'.encode(), hashlib.sha256).hexdigest()[:32], sig):
                return None
            if int(exp) < datetime.now().timestamp():
                return None
            return user, None
        user, ver, exp, sig = parts
        if not hmac.compare_digest(
                hmac.new(_token_secret(), f'{user}|{ver}|{exp}'.encode(), hashlib.sha256).hexdigest()[:32], sig):
            return None
        if int(exp) < datetime.now().timestamp():
            return None
        return user, int(ver)
    except Exception:
        return None


def b64e(raw: bytes) -> str:
    return base64.urlsafe_b64encode(raw).decode().rstrip('=')


def b64d(tok: str) -> bytes:
    return base64.urlsafe_b64decode(tok + '=' * (-len(tok) % 4))


def auth_user():
    """当前登录用户：优先 session，其次 X-Sup-Token / ?token= 。→ (user, role) 或 (None, None)"""
    if session.get('user'):
        return session['user'], session.get('role')
    tok = request.headers.get('X-Sup-Token') or request.args.get('token') or ''
    if not tok:
        return None, None
    parsed = verify_token(tok)
    if not parsed:
        return None, None
    user, ver = parsed
    con = db()
    row = con.execute('SELECT role, pwd_ver FROM users WHERE username=?', (user,)).fetchone()
    con.close()
    if not row:
        return None, None
    if ver is not None and int(row['pwd_ver'] or 1) != ver:   # 改过密码 → 旧令牌失效（旧格式令牌不做版本校验）
        return None, None
    return user, row['role']


def login_required(f):
    @wraps(f)
    def wrap(*a, **kw):
        if not auth_user()[0]:
            return jsonify(error='未登录'), 401
        return f(*a, **kw)
    return wrap


def admin_required(f):
    @wraps(f)
    def wrap(*a, **kw):
        user, role = auth_user()
        if not user:
            return jsonify(error='未登录'), 401
        if role != 'admin':
            return jsonify(error='需要管理员权限'), 403
        return f(*a, **kw)
    return wrap


# ────────────────────────── 页面路由 ──────────────────────────

@bp.route('/')
def index():
    if not session.get('user'):
        return render_template('index.html', user=None)
    return render_template('index.html', user=session['user'], role=session.get('role'))


@bp.route('/admin')
def admin_page():
    if not session.get('user'):
        # 未登录：带着 next=admin 回登录页，登录后由前端送回 /admin（否则会停在 App 首页）
        return redirect(url_for('sup.index', next='admin'))
    if session.get('role') != 'admin':
        return redirect(url_for('sup.index'))     # 已登录但非管理员：不提示、直接回 App
    return render_template('admin.html', user=session['user'])


# ────────────────────────── 认证 API ──────────────────────────

@bp.route('/api/login', methods=['POST'])
def api_login():
    data = request.get_json(force=True)
    username = (data.get('username') or '').strip()
    password = data.get('password') or ''
    con = db()
    row = con.execute('SELECT * FROM users WHERE username=?', (username,)).fetchone()
    con.close()
    if not row or row['password_hash'] != hashlib.sha256(password.encode()).hexdigest():
        return jsonify(error='用户名或密码错误'), 401
    session.permanent = True          # 30 天持久会话（见 create_app 注释）
    session['user'] = row['username']
    session['role'] = row['role']
    # 返回 user：登录是 fetch 静默完成（页面不刷新），前端需回写 #whoami 的 data-user，
    # 否则验收联动/{{拍照人}} 拿到的用户名为空（v0.8.2 修复）
    return jsonify(ok=True, role=row['role'], user=row['username'],
                   token=make_token(row['username'], row['pwd_ver'] or 1))


def _set_password(con, username: str, new_pwd: str) -> int:
    """写入新密码哈希并把 pwd_ver +1（旧令牌随之失效）；返回新的 pwd_ver。"""
    con.execute('UPDATE users SET password_hash=?, pwd_ver=pwd_ver+1 WHERE username=?',
                (hashlib.sha256(new_pwd.encode()).hexdigest(), username))
    row = con.execute('SELECT pwd_ver, role FROM users WHERE username=?', (username,)).fetchone()
    con.commit()
    return int(row['pwd_ver'] or 1)


def _pwd_error(new_pwd: str):
    if not new_pwd or len(new_pwd) < 4:
        return '新密码至少 4 位'
    if len(new_pwd) > 64:
        return '新密码过长'
    return None


@bp.route('/api/password', methods=['POST'])
@login_required
def api_change_password():
    """本人修改自己的密码（v0.21）：需验证原密码；成功后返回新令牌（当前设备免重登）。"""
    user, _role = auth_user()
    data = request.get_json(force=True) or {}
    old_pwd = data.get('old_password') or ''
    new_pwd = (data.get('new_password') or '').strip()
    err = _pwd_error(new_pwd)
    if err:
        return jsonify(error=err), 400
    con = db()
    try:
        row = con.execute('SELECT password_hash FROM users WHERE username=?', (user,)).fetchone()
        if not row or row['password_hash'] != hashlib.sha256(old_pwd.encode()).hexdigest():
            return jsonify(error='原密码不正确'), 403
        ver = _set_password(con, user, new_pwd)
        return jsonify(ok=True, token=make_token(user, ver))
    finally:
        con.close()


@bp.route('/admin/api/users', methods=['GET'])
@admin_required
def admin_users():
    """用户列表（管理员）：姓名 / 角色（不返回任何密码信息）。"""
    con = db()
    rows = con.execute('SELECT username, role FROM users ORDER BY role DESC, username').fetchall()
    con.close()
    return jsonify(users=[{'username': r['username'], 'role': r['role']} for r in rows])


@bp.route('/admin/api/users/<path:username>/password', methods=['POST'])
@admin_required
def admin_set_password(username):
    """管理员为他人重置密码（无需原密码）；被改者旧令牌立即失效，需用新密码重新登录。"""
    new_pwd = ((request.get_json(force=True) or {}).get('new_password') or '').strip()
    err = _pwd_error(new_pwd)
    if err:
        return jsonify(error=err), 400
    con = db()
    try:
        if not con.execute('SELECT 1 FROM users WHERE username=?', (username,)).fetchone():
            return jsonify(error=f'用户不存在：{username}'), 404
        _set_password(con, username, new_pwd)
        return jsonify(ok=True, username=username)
    finally:
        con.close()


@bp.route('/api/me', methods=['GET'])
def api_me():
    """当前登录身份（页面重载后用令牌恢复用户名，供 {{拍照人}} / 验收联动使用）。"""
    user, role = auth_user()
    if not user:
        return jsonify(error='未登录'), 401
    return jsonify(user=user, role=role)


@bp.route('/api/logout', methods=['POST'])
def api_logout():
    session.clear()
    return jsonify(ok=True)


# ────────────────────────── 工作簿 API ──────────────────────────

def _norm_cell(v):
    """单元格归一化。

    注意：**必须先判 datetime 再判 date** —— datetime 是 date 的子类，
    且 str(datetime) 会得到 '2026-09-15 00:00:00'，会把日期列写成带时间的字符串
    （导出的 Excel 里日期是真日期单元格，回传时必须还原成 'YYYY-MM-DD'，否则数据被污染）。
    """
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d') if v.time() == time.min else v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, (int, float)):
        return v
    return str(v).strip()


def _norm_rows(headers, rows, config=None):
    """株数强度/蓄积强度等：数值 0.18 存储统一规范化为文本 '18%'（doc/004 §二）。

    v0.23（C11）：转换哪些列由参数「百分比列」声明；未声明回落旧规则
    「列名含『强度』且不含 %」，故旧模板行为完全不变。
    """
    pct_idx = percent_cols_of(headers, config or {})
    out = []
    for r in rows:
        r = [_norm_cell(v) for v in r]
        for i in pct_idx:
            v = r[i]
            if isinstance(v, (int, float)) and not isinstance(v, bool) and 0 < v < 1:
                r[i] = f'{v * 100:g}%'
        out.append(r)
    return out


PARAM_SHEET = '参数'   # 参数 sheet 固定名（C02）
DATA_SHEET = 'data'    # 数据 sheet 约定名（v0.23，C11）


def _pick_data_sheet(sheetnames):
    """定位数据 sheet（v0.23，C11）。

    优先精确取约定名「data」；没有则要求"有且仅有 1 个非参数 sheet"；
    多个非参数 sheet 时**明确报错**——旧写法 `next(n for n in sheetnames if n != '参数')`
    会静默取第一个，工作簿一带说明页/汇总页就取错表。
    """
    others = [n for n in sheetnames if n != PARAM_SHEET]
    if not others:
        raise ParamError(f'上传的 Excel 只有「{PARAM_SHEET}」sheet，没有数据 sheet')
    if DATA_SHEET in others:
        return DATA_SHEET
    if len(others) > 1:
        raise ParamError(
            f'无法判断用哪个数据 sheet：共有 {len(others)} 个非「{PARAM_SHEET}」sheet {others}。'
            f'请把数据 sheet 命名为「{DATA_SHEET}」，或只保留一个（其余说明/汇总页请删掉）')
    return others[0]


def parse_workbook_storage(stream):
    """解析上传的 xlsx：数据 sheet + 参数 sheet。违规抛 ParamError。

    返回 (sheet_name, headers, rows, config, param_rows, key_column)。
    声明了「unique-key」时对唯一键做非空+唯一校验，不过则拒绝上传（C11）。
    """
    wb = openpyxl.load_workbook(stream, data_only=True, read_only=True)
    try:
        if PARAM_SHEET not in wb.sheetnames:
            raise ParamError(
                f'上传的 Excel 缺少「{PARAM_SHEET}」sheet（需同时含数据 sheet 和「{PARAM_SHEET}」sheet；'
                f'当前 sheet：{wb.sheetnames}）')
        data_name = _pick_data_sheet(wb.sheetnames)

        ws_p = wb[PARAM_SHEET]
        prows = [(i, r) for i, r in enumerate(
            ws_p.iter_rows(max_col=6, values_only=True), 1)
            if any(x not in (None, '') for x in r)]

        ws_d = wb[data_name]
        it = ws_d.iter_rows(max_col=64, values_only=True)  # read_only 维度可能误报列数，强制扩列
        first = next(it, None)
        if first is None:
            raise ParamError(f'数据 sheet「{data_name}」为空')
        headers = [_s(h) for h in first]
        while headers and not headers[-1]:
            headers.pop()
        if not headers:
            raise ParamError(f'数据 sheet「{data_name}」缺少表头行')
        ncols = len(headers)
        rows = []
        for r in it:
            vals = list(r) + [None] * (ncols - len(r))
            if any(v not in (None, '') for v in vals[:ncols]):
                rows.append(vals[:ncols])

        config = parse_params(prows, headers)
        rows = _norm_rows(headers, rows, config)
        # C11：声明了 unique-key 才校验（未声明 = 不校验，旧模板零改动）
        uerr = check_unique_column(rows, headers, config)
        if uerr:
            raise ParamError(f'数据 sheet「{data_name}」：{uerr}')
        param_rows = [[_norm_cell(v) for v in (list(r) + [None] * 6)[:6]] for _, r in prows]
        return data_name, headers, rows, config, param_rows, row_key_column(config)
    finally:
        wb.close()


@bp.route('/api/workbooks', methods=['GET'])
@login_required
def api_list():
    """App 端工作簿列表：**只列未下架的**（C13）——下架的由后台「显示已下架」查看。"""
    con = db()
    try:
        rows = con.execute(
            'SELECT id, name, sheet_name, uploaded_at FROM workbooks'
            ' WHERE COALESCE(is_active, 1)=1 ORDER BY id DESC').fetchall()
    except sqlite3.OperationalError:      # 老库未迁移（无 is_active 列）
        rows = con.execute(
            'SELECT id, name, sheet_name, uploaded_at FROM workbooks ORDER BY id DESC').fetchall()
    con.close()
    return jsonify(workbooks=[dict(r) for r in rows])


@bp.route('/api/workbooks', methods=['POST'])
@login_required
def api_upload():
    fs = request.files.get('file')
    if not fs or not fs.filename.lower().endswith(('.xlsx', '.xlsm')):
        return jsonify(error='请上传 .xlsx 文件'), 400
    blob = fs.read()   # 先读全量字节：既供解析，也原样留存作导出模板
    try:
        (sheet_name, headers, rows, config,
         param_rows, key_column) = parse_workbook_storage(io.BytesIO(blob))
    except ParamError as e:
        return jsonify(error=str(e)), 400
    con = db()
    cur = con.execute(
        'INSERT INTO workbooks(name, sheet_name, headers, rows, config, param_rows,'
        ' key_column, uploaded_at) VALUES(?,?,?,?,?,?,?,?)',
        (fs.filename, sheet_name, json.dumps(headers, ensure_ascii=False),
         json.dumps(rows, ensure_ascii=False), json.dumps(config, ensure_ascii=False),
         json.dumps(param_rows, ensure_ascii=False), key_column,
         datetime.now().strftime('%Y-%m-%d %H:%M')))
    wid = cur.lastrowid
    con.commit()
    con.close()
    # 原始模板留存（导出按上传模板回填，保留全部格式）
    src_dir = UPLOAD_DIR / str(wid)
    src_dir.mkdir(parents=True, exist_ok=True)
    ext = '.xlsm' if fs.filename.lower().endswith('.xlsm') else '.xlsx'
    (src_dir / f'source{ext}').write_bytes(blob)
    return jsonify(ok=True, id=wid, rows=len(rows))


def _wb_key_column(r):
    """工作簿的生效唯一键列名（C11）：落库值优先，为空则从 config 推。

    老工作簿（v0.23 之前上传的）key_column 为 ''，这里回落到「小班号」，行为不变。
    """
    try:
        kc = r['key_column']
    except (IndexError, KeyError):
        kc = ''
    if kc:
        return kc
    try:
        cfg = json.loads(r['config'])
    except (TypeError, ValueError):
        cfg = {}
    return row_key_column(cfg)


# ── 下架 / 上架（C13）──────────────────────────────────────────────
# 语义：下架 = **App/user 端看不见、不能操作**，但数据行、源模板、日志、照片全部保留，
# 管理员随时可「上架」恢复。它是**软删除**，不是权限——用户端的按钮是否显示与它无关，
# 真正的边界仍是服务端的 @admin_required。
INACTIVE_MSG = '该工作簿已被管理员下架，App 端暂不可用（数据已保留，请联系管理员恢复）'


def _inactive_for(con, wid):
    """该工作簿对**当前用户**是否不可用（已下架且当前用户非管理员）→ True/False。"""
    try:
        r = con.execute('SELECT COALESCE(is_active, 1) AS a FROM workbooks WHERE id=?',
                        (wid,)).fetchone()
    except sqlite3.OperationalError:
        return False                      # 老库未迁移 → 视为全部在用
    if not r:
        return False                      # 不存在交给各路由自己 404
    # 注意：不能写 `r['a'] or 1` —— is_active=0（已下架）会被 or 吃成 1，
    # 导致守卫永远认为"在用"（2026-09-14 实测踩坑）。
    if r['a'] is None or int(r['a']) == 1:
        return False
    _u, role = auth_user()
    return role != 'admin'


@bp.route('/api/workbooks/<int:wid>', methods=['GET'])
@login_required
def api_get(wid):
    con = db()
    r = con.execute('SELECT * FROM workbooks WHERE id=?', (wid,)).fetchone()
    inactive = _inactive_for(con, wid)
    con.close()
    if not r:
        abort(404)
    if inactive:
        return jsonify(error=INACTIVE_MSG), 403
    headers = json.loads(r['headers'])
    cfg = json.loads(r['config'])
    return jsonify(id=r['id'], name=r['name'], sheet_name=r['sheet_name'],
                   headers=headers, rows=json.loads(r['rows']),
                   config=cfg, key_column=_wb_key_column(r),
                   # v0.23：审计/联动字段由后端算好给前端，避免前后端各写一份默认值
                   log_fields=log_fields_of(headers, cfg))


@bp.route('/api/workbooks/<int:wid>', methods=['DELETE'])
@admin_required
def api_delete(wid):
    """**下架**（软删除，C13）：is_active=0 —— App/user 端不再显示、不可操作，
    数据行 / 源模板 / 日志 / 照片全部保留，可用「上架」恢复。

    `?purge=1` 才是**彻底删除**：连 DB 行、源模板目录、accept_logs/photo_sync 一起清。
    （v0.23 之前 DELETE 就是彻底删但不清目录，会留孤儿目录——已修。）
    """
    purge = request.args.get('purge') == '1'
    con = db()
    if purge:
        con.execute('DELETE FROM workbooks WHERE id=?', (wid,))
        con.execute('DELETE FROM accept_logs WHERE workbook_id=?', (wid,))
        con.execute('DELETE FROM photo_sync WHERE workbook_id=?', (wid,))
        con.commit()
        con.close()
        shutil.rmtree(UPLOAD_DIR / str(wid), ignore_errors=True)   # 源模板目录
    else:
        cur = con.execute('UPDATE workbooks SET is_active=0 WHERE id=?', (wid,))
        n = cur.rowcount
        con.commit()
        con.close()
        if not n:
            abort(404)
    return jsonify(ok=True, purged=purge)


@bp.route('/api/workbooks/<int:wid>/restore', methods=['POST'])
@admin_required
def api_restore(wid):
    """上架：撤销下架，App 端恢复可见可用（C13）。"""
    con = db()
    cur = con.execute('UPDATE workbooks SET is_active=1 WHERE id=?', (wid,))
    n = cur.rowcount
    con.commit()
    con.close()
    if not n:
        abort(404)
    return jsonify(ok=True)


# ── 更新 Excel（按唯一键合并，v0.26）────────────────────────────────
# 场景：工作簿已经填了一半，此时要改模板（加「导出筛选」、改「目录」、加列、修正原始数据）。
# 原做法只能重新上传 → 会**新建**一个工作簿（已填数据成了孤儿，变更日志/下架状态全断链）。
# 这里改为**就地更新**：保留原 id，把新 Excel 按唯一键合并进现有数据行。
MAX_BACKUPS = 5          # 每个工作簿最多保留几份"更新前"备份


@bp.route('/api/workbooks/<int:wid>/update', methods=['POST'])
@admin_required
def api_update(wid):
    """按唯一键**就地更新**已有工作簿（保留原 id）。

    - 以既有行顺序为基准原地合并 → 老行的 row_idx 不变，变更日志不错位；新表的行追加末尾
    - 冲突：人工填过的列（可编辑列，旧∪新）旧值非空优先；**其余列以新 Excel 为准**
    - `?dry=1` 只算不写、返回统计供确认；正式提交需 `confirm=<当前工作簿名>`（C13 的确认规矩）
    - 写库前自动备份（DB 行 + 源模板文件），可用 `/rollback` 还原
    """
    fs = request.files.get('file')
    if not fs or not fs.filename.lower().endswith(('.xlsx', '.xlsm')):
        return jsonify(error='请上传 .xlsx 文件'), 400
    blob = fs.read()
    con = db()
    try:
        r = con.execute('SELECT * FROM workbooks WHERE id=?', (wid,)).fetchone()
        if not r:
            abort(404)
        try:
            (new_sheet, new_headers, new_rows,
             new_cfg, new_param_rows, new_key) = parse_workbook_storage(io.BytesIO(blob))
        except ParamError as e:
            return jsonify(error=f'新 Excel 校验未通过：{e}'), 400

        old_headers = json.loads(r['headers'])
        old_rows = json.loads(r['rows'])
        old_cfg = json.loads(r['config'])
        old_key = _wb_key_column(r)
        if new_key != old_key:
            return jsonify(error=(
                f'唯一键不一致：现有工作簿按「{old_key}」定位行，新 Excel 按「{new_key}」。'
                f'两边 unique-key 需要一致才能按行更新')), 400

        # 人工在 App 里能填的列 = 旧表「可编辑列」∪ 新表「可编辑列」（旧表填过的新表删了也要保住）
        preserve = sorted(set(split_list(old_cfg.get('可编辑列', ''))) |
                          set(split_list(new_cfg.get('可编辑列', ''))))
        try:
            merged, stats = merge_rows(old_headers, old_rows,
                                      new_headers, new_rows, old_key, preserve)
        except ParamError as e:
            return jsonify(error=str(e)), 400

        dry = request.args.get('dry') == '1'
        stats.update({
            'workbook_id': wid, 'name': r['name'], 'new_name': fs.filename,
            'old_row_count': len(old_rows), 'new_row_count': len(new_rows),
            'result_row_count': len(merged),
            'old_col_count': len(old_headers), 'new_col_count': len(new_headers),
            'dry': dry,
        })
        if dry:
            return jsonify(ok=True, preview=stats)

        confirm = (request.form.get('confirm') or '').strip()
        if confirm != r['name']:
            return jsonify(error='确认名称不一致：请照原样输入当前工作簿名称'), 400

        # ① 备份（DB 行 + 源模板文件）
        ext = '.xlsm' if r['name'].lower().endswith('.xlsm') else '.xlsx'
        d = UPLOAD_DIR / str(wid)
        d.mkdir(parents=True, exist_ok=True)
        src_bak = f'source_bak_{datetime.now().strftime("%Y%m%d%H%M%S")}{ext}'
        if (d / f'source{ext}').exists():
            shutil.copy2(d / f'source{ext}', d / src_bak)
        cur = con.execute(
            'INSERT INTO workbook_backups(workbook_id,name,sheet_name,headers,rows,config,'
            ' param_rows,key_column,src_ext,src_bak,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',
            (wid, r['name'], r['sheet_name'], r['headers'], r['rows'], r['config'],
             r['param_rows'], r['key_column'], ext, src_bak,
             datetime.now().strftime('%Y-%m-%d %H:%M:%S')))
        bid = cur.lastrowid
        for ob in con.execute('SELECT id, src_bak FROM workbook_backups WHERE workbook_id=?'
                              ' ORDER BY id DESC', (wid,)).fetchall()[MAX_BACKUPS:]:
            con.execute('DELETE FROM workbook_backups WHERE id=?', (ob['id'],))
            if ob['src_bak']:
                (d / ob['src_bak']).unlink(missing_ok=True)

        # ② 就地覆盖（id / is_active 不变 → 变更日志、下架状态、App 使用习惯全部保住）
        new_ext = '.xlsm' if fs.filename.lower().endswith('.xlsm') else '.xlsx'
        con.execute(
            'UPDATE workbooks SET name=?, sheet_name=?, headers=?, rows=?, config=?,'
            ' param_rows=?, key_column=? WHERE id=?',
            (fs.filename, new_sheet, json.dumps(new_headers, ensure_ascii=False),
             json.dumps(merged, ensure_ascii=False), json.dumps(new_cfg, ensure_ascii=False),
             json.dumps(new_param_rows, ensure_ascii=False), new_key, wid))
        con.commit()

        # ③ 换成新的源模板（导出按新模板回填）；扩展名变了就清掉旧的
        (d / f'source{new_ext}').write_bytes(blob)
        for other in ('.xlsx', '.xlsm'):
            if other != new_ext:
                (d / f'source{other}').unlink(missing_ok=True)

        stats.update({'backup_id': bid, 'src_bak': src_bak})
        return jsonify(ok=True, updated=True, preview=stats)
    finally:
        con.close()


@bp.route('/api/workbooks/<int:wid>/rollback', methods=['POST'])
@admin_required
def api_rollback(wid):
    """回滚到最近一次「更新 Excel」之前（数据 + 参数 + 源模板一起还原）。"""
    data = request.get_json(silent=True) or {}
    con = db()
    try:
        r = con.execute('SELECT id, name FROM workbooks WHERE id=?', (wid,)).fetchone()
        if not r:
            abort(404)
        b = con.execute('SELECT * FROM workbook_backups WHERE workbook_id=?'
                        ' ORDER BY id DESC LIMIT 1', (wid,)).fetchone()
        if not b:
            return jsonify(error='没有可回滚的备份（只有执行过「更新 Excel」才会有）'), 400
        if (data.get('confirm') or '').strip() != r['name']:
            return jsonify(error='确认名称不一致：请照原样输入当前工作簿名称'), 400

        con.execute(
            'UPDATE workbooks SET name=?, sheet_name=?, headers=?, rows=?, config=?,'
            ' param_rows=?, key_column=? WHERE id=?',
            (b['name'], b['sheet_name'], b['headers'], b['rows'], b['config'],
             b['param_rows'], b['key_column'], wid))
        con.execute('DELETE FROM workbook_backups WHERE id=?', (b['id'],))
        con.commit()

        d = UPLOAD_DIR / str(wid)
        bak = (d / b['src_bak']) if b['src_bak'] else None
        if bak and bak.exists():
            for other in ('.xlsx', '.xlsm'):
                if other != b['src_ext']:
                    (d / f'source{other}').unlink(missing_ok=True)
            shutil.copy2(bak, d / f'source{b["src_ext"]}')
            bak.unlink(missing_ok=True)
        return jsonify(ok=True, rolled_back=True,
                       restored={'name': b['name'], 'rows': len(json.loads(b['rows']))})
    finally:
        con.close()


# 变更日志缺省审计字段（v0.23，C11）：列名在模板中存在才记录。
# 缺省值现由 param_parser.DEFAULT_LOG_FIELDS 单一维护（验收人/验收日期/验收时间/验收结果/验收备注），
# 要改审计范围时在「参数」sheet 加一行 key=日志字段 即可，不必改代码。


@bp.route('/api/workbooks/<int:wid>/rows/<int:ridx>', methods=['POST'])
@login_required
def api_save_row(wid, ridx):
    """单行自动保存（前端两列表单 onchange 触发，对齐 hqz-survey 编辑体验）。
    审计字段的每次变化写入 accept_logs（操作人+时间）。"""
    vals = (request.get_json(silent=True) or {}).get('values')
    if not isinstance(vals, list):
        return jsonify(error='values 必须为数组'), 400
    con = db()
    try:
        r = con.execute('SELECT headers, rows, config, key_column FROM workbooks WHERE id=?',
                        (wid,)).fetchone()
        if not r:
            abort(404)
        if _inactive_for(con, wid):        # C13：下架后 user 端不可操作
            return jsonify(error=INACTIVE_MSG), 403
        headers = json.loads(r['headers'])
        rows = json.loads(r['rows'])
        cfg = json.loads(r['config'])
        if not (0 <= ridx < len(rows)):
            abort(404)
        if len(vals) != len(rows[ridx]):
            return jsonify(error=f'列数不匹配：期望 {len(rows[ridx])} 列，收到 {len(vals)}'), 400
        old_row = rows[ridx]
        rows[ridx] = [_norm_cell(v) for v in vals]

        # 审计字段变更日志（v0.23：列清单由「日志字段」参数驱动）
        log_fields = set(log_fields_of(headers, cfg))
        logs = []
        for i, h in enumerate(headers):
            if h not in log_fields:
                continue
            ov = '' if old_row[i] is None else str(old_row[i]).strip()
            nv = '' if rows[ridx][i] is None else str(rows[ridx][i]).strip()
            if ov != nv:
                logs.append((h, ov, nv))
        if logs:
            # 行标识取自工作簿的唯一键列（C11），不再硬编码「小班号」
            kc = _wb_key_column(r)
            hi = headers.index(kc) if kc in headers else -1
            xiaoban = '' if hi < 0 else str(old_row[hi] or '').strip()
            now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            op = session.get('user') or ''
            con.executemany(
                'INSERT INTO accept_logs(workbook_id,row_idx,xiaoban,field,old_value,new_value,operator,created_at)'
                ' VALUES(?,?,?,?,?,?,?,?)',
                [(wid, ridx, xiaoban, f, ov, nv, op, now) for f, ov, nv in logs])

        con.execute('UPDATE workbooks SET rows=? WHERE id=?',
                    (json.dumps(rows, ensure_ascii=False), wid))
        con.commit()
    finally:
        con.close()
    return jsonify(ok=True)


@bp.route('/api/workbooks/<int:wid>/save', methods=['POST'])
@login_required
def api_save(wid):
    rows = request.get_json(force=True).get('rows')
    if rows is None:
        return jsonify(error='缺少 rows'), 400
    con = db()
    if _inactive_for(con, wid):            # C13：下架后 user 端不可操作
        con.close()
        return jsonify(error=INACTIVE_MSG), 403
    cur = con.execute('UPDATE workbooks SET rows=? WHERE id=?',
                      (json.dumps(rows, ensure_ascii=False), wid))
    con.commit()
    n = cur.rowcount
    con.close()
    if not n:
        abort(404)
    return jsonify(ok=True)


# ────────────────────────── 轨迹（B3 前端接入） ──────────────────────────

@bp.route('/api/track', methods=['POST'])
@login_required
def api_track():
    fs = request.files.get('file')
    if not fs or not fs.filename.lower().endswith('.gpx'):
        return jsonify(error='请上传 .gpx 轨迹文件'), 400
    safe = re.sub(r'[\\/:*?"<>|]+', '_', fs.filename)
    path = TRACK_DIR / f'{datetime.now().strftime("%Y%m%d_%H%M%S")}_{safe}'
    fs.save(path)
    return jsonify(ok=True, file=path.name)


# ────────────────────────── 相片云同步（doc/007） ──────────────────────────
# C07 修订：服务器不长期存储——照片只在七牛 zz-1 暂存，百度成功后按保留期清理。
# 仅 App 端调用（浏览器端不拍照）；开关关闭 → 423，前端静默跳过。

_PHOTO_MAX = 2 * 1024 * 1024  # 2MB（App 端压缩后 300-500KB）


@bp.route('/api/sync/enabled', methods=['GET'])
@login_required
def api_sync_enabled():
    con = db()
    s = sync_cloud.get_settings(con)
    con.close()
    return jsonify(enabled=s.get('sync_enabled') == '1')


@bp.route('/api/photo', methods=['POST'])
@login_required
def api_photo():
    con = db()
    try:
        s = sync_cloud.get_settings(con)
        if s.get('sync_enabled') != '1':
            return jsonify(error='云同步未开启（后台「同步设置」）'), 423
        fs = request.files.get('file')
        if not fs:
            return jsonify(error='缺少 file'), 400
        data = fs.read()
        if not data or len(data) > _PHOTO_MAX:
            return jsonify(error=f'照片大小超出限制（≤{_PHOTO_MAX // 1024}KB）'), 400
        r = request.form
        try:
            wid = int(r.get('workbook_id') or 0)
        except ValueError:
            return jsonify(error='workbook_id 无效'), 400
        if wid and _inactive_for(con, wid):    # C13：下架后不再接收该工作簿的照片
            return jsonify(error=INACTIVE_MSG), 403
        filename = re.sub(r'[\\/:*?"<>|]+', '_', (r.get('filename') or '').strip())
        if not filename.lower().endswith('.jpg'):
            filename += '.jpg'
        subdir = re.sub(r'^/+|/+$', '', r.get('subdir') or '')
        if not all(re.match(r'^[^\\/:*?"<>|]+$', seg) for seg in subdir.split('/') if seg):
            return jsonify(error='subdir 含非法字符'), 400
        xiaoban = (r.get('xiaoban') or '').strip()[:64]

        qiniu_key, remote_path = sync_cloud.photo_paths(s, subdir, filename)
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        cur = con.execute(
            'INSERT INTO photo_sync(workbook_id, xiaoban, filename, remote_path,'
            ' qiniu_key, size, state, created_at) VALUES(?,?,?,?,?,?,?,?)',
            (wid, xiaoban, filename, remote_path, qiniu_key, len(data), 'received', now))
        pid = cur.lastrowid
        con.commit()
        # 待推送缓冲（data/pending/{id}.jpg）：百度成功即删（C07 修订：不长期存储）
        (PENDING_DIR / f'{pid}.jpg').write_bytes(data)

        # 第一跳：七牛暂存（异地备份；失败不阻塞——本地缓冲可支撑后续推送）
        try:
            sync_cloud.QiniuClient(s['sync_qiniu_ak'], s['sync_qiniu_sk'],
                                   s['sync_qiniu_bucket']).put_bytes(
                qiniu_key, data, mime='image/jpeg')
            con.execute("UPDATE photo_sync SET state='qiniu_ok' WHERE id=?", (pid,))
            con.commit()
        except sync_cloud.SyncError as e:
            con.execute('UPDATE photo_sync SET retry_count=retry_count+1, last_error=?'
                        ' WHERE id=?', (str(e)[:500], pid))
            con.commit()

        # 第二跳（懒触发）：立即尝试一次百度推送，失败留给积压等后台手动重推
        pushed = False
        try:
            row = con.execute('SELECT * FROM photo_sync WHERE id=?', (pid,)).fetchone()
            sync_cloud.push_to_baidu(con, row, PENDING_DIR)
            con.commit()
            pushed = True
        except sync_cloud.SyncError as e:
            con.execute("UPDATE photo_sync SET retry_count=retry_count+1, last_error=?"
                        " WHERE id=?", (str(e)[:500], pid))
            con.commit()
        # 积压补推 + 七牛副本清理：放后台 daemon 线程（不阻塞拍照响应；
        # 线程用独立 sqlite 连接，gunicorn worker 退出被杀也无害——下次触发自动续做）
        if pushed:
            def _bg_sync():
                try:
                    c = db()
                    sync_cloud.try_push_pending(c, limit=2, pending_dir=PENDING_DIR)
                    sync_cloud.purge_expired(c, limit=20)
                    c.close()
                except Exception:
                    pass
            threading.Thread(target=_bg_sync, daemon=True).start()
        row = con.execute('SELECT state, last_error FROM photo_sync WHERE id=?', (pid,)).fetchone()
        return jsonify(ok=True, id=pid, state=row['state'],
                       error=row['last_error'] or None)
    finally:
        con.close()


# ────────────────────────── 视频（v0.10）：服务器压缩 → 七牛 → 百度 ──────────────────────────
# 浏览器/WebView 无法转码视频，压缩统一在服务器用 ffmpeg 完成（参数见后台「视频压缩」）。
# 视频体积大：接收即返回（state=received），压缩+推送在后台线程进行，不阻塞手机。

def _video_name(base: str) -> str:
    """视频文件名 = 相片文件名模板渲染结果 + `_视频` 标记（与照片命名区分）。"""
    name = re.sub(r'[\\/:*?"<>|]+', '_', (base or '').strip()) or 'video'
    if not name.endswith('_视频'):
        name += '_视频'
    return name


def _process_video(pid: int, skip_compress: bool = False) -> None:
    """后台：ffmpeg 压缩（手机端已压缩则跳过）→ 七牛暂存 → 百度网盘；失败留状态待重推。"""
    con = db()
    try:
        row = con.execute('SELECT * FROM photo_sync WHERE id=?', (pid,)).fetchone()
        if not row:
            return
        s = sync_cloud.get_settings(con)
        raw = PENDING_DIR / f'{pid}.mp4'
        comp = PENDING_DIR / f'{pid}_c.mp4'
        if not raw.exists():
            con.execute('UPDATE photo_sync SET last_error=? WHERE id=?',
                        ('原始视频缓冲丢失（无法重推）', pid))
            con.commit()
            return
        if skip_compress:
            src = raw  # 手机端已压缩（MediaRecorder），服务器不再重复转码
            con.execute('UPDATE photo_sync SET size=? WHERE id=?', (raw.stat().st_size, pid))
        else:
            ok, msg = sync_cloud.compress_video(raw, comp, s)
            src = comp if ok else raw
            if ok:
                con.execute('UPDATE photo_sync SET size=? WHERE id=?', (comp.stat().st_size, pid))
            else:
                con.execute('UPDATE photo_sync SET last_error=? WHERE id=?',
                            (f'压缩失败（改用原文件）：{msg}'[:400], pid))
        con.commit()
        # 七牛暂存（异地备份）
        try:
            mime = 'video/mp4' if row['qiniu_key'].lower().endswith('.mp4') else 'video/webm'
            sync_cloud.QiniuClient(s['sync_qiniu_ak'], s['sync_qiniu_sk'],
                                   s['sync_qiniu_bucket']).put_bytes(
                row['qiniu_key'], src.read_bytes(), mime=mime)
            con.execute("UPDATE photo_sync SET state='qiniu_ok' WHERE id=?", (pid,))
            con.commit()
        except sync_cloud.SyncError as e:
            con.execute('UPDATE photo_sync SET retry_count=retry_count+1, last_error=? WHERE id=?',
                        (str(e)[:500], pid))
            con.commit()
        # 百度推送
        row = con.execute('SELECT * FROM photo_sync WHERE id=?', (pid,)).fetchone()
        try:
            sync_cloud.push_to_baidu(con, row, PENDING_DIR)
            con.commit()
            sync_cloud.try_push_pending(con, limit=1, pending_dir=PENDING_DIR)
            sync_cloud.purge_expired(con, limit=20)
        except sync_cloud.SyncError as e:
            con.execute('UPDATE photo_sync SET retry_count=retry_count+1, last_error=? WHERE id=?',
                        (str(e)[:500], pid))
            con.commit()
    except Exception:
        pass
    finally:
        con.close()


def _num(d, k, default, cast=float):
    try:
        return cast(d.get(k) or default)
    except (ValueError, TypeError):
        return default


def _video_params(s: dict) -> dict:
    return {
        'rec': s.get('video_phone_rec', '1'),
        'rec_mode': (s.get('video_rec_mode') or 'system'),
        'transcode': (s.get('video_transcode') or '1'),
        'cam_quality': _num(s, 'video_cam_quality', 1, int),
        'max_height': _num(s, 'video_max_height', 1080, int),
        'bitrate_k': _num(s, 'video_maxrate_k', 4000, int),
        'fps': _num(s, 'video_fps', 30, int),
        'max_seconds': _num(s, 'video_max_seconds', 60, int),
        'max_mb': _num(s, 'video_max_mb', 300, int),
    }


def _photo_params(s: dict) -> dict:
    """图片压缩服务器级默认（工作簿参数 sheet 优先于这里）。"""
    return {
        'max_side': _num(s, 'photo_max_side', 1440, int),
        'quality': round(min(1.0, max(0.1, _num(s, 'photo_quality', 0.8))), 3),
    }


@bp.route('/api/compress/params', methods=['GET'])
@login_required
def api_compress_params():
    """压缩参数（只读）：photo=图片压缩默认，video=视频录制/转码参数。

    与「同步」彻底解耦：压缩在手机端/本地完成，是否上云由 sync_enabled 决定。
    """
    con = db()
    try:
        s = sync_cloud.get_settings(con)
        return jsonify(photo=_photo_params(s), video=_video_params(s))
    finally:
        con.close()


@bp.route('/api/video/params', methods=['GET'])
@login_required
def api_video_params():
    """手机端录制压缩参数（保留旧端点以兼容旧包；等价于 /api/compress/params 的 video 部分）。"""
    con = db()
    try:
        return jsonify(**_video_params(sync_cloud.get_settings(con)))
    finally:
        con.close()


@bp.route('/api/video', methods=['POST'])
@login_required
def api_video():
    con = db()
    try:
        s = sync_cloud.get_settings(con)
        if s.get('sync_enabled') != '1':
            return jsonify(error='视频需先开启云同步（后台「同步设置」）'), 423
        fs = request.files.get('file')
        if not fs:
            return jsonify(error='缺少 file'), 400
        try:
            max_mb = int(s.get('video_max_mb') or 300)
        except ValueError:
            max_mb = 300
        data = fs.read()
        if not data:
            return jsonify(error='视频为空'), 400
        if len(data) > max_mb * 1024 * 1024:
            return jsonify(error=f'视频超出上限 {max_mb}MB（后台「视频压缩」可调）'), 400
        r = request.form
        try:
            wid = int(r.get('workbook_id') or 0)
        except ValueError:
            return jsonify(error='workbook_id 无效'), 400
        if wid and _inactive_for(con, wid):    # C13：下架后不再接收该工作簿的视频
            return jsonify(error=INACTIVE_MSG), 403
        subdir = re.sub(r'^/+|/+$', '', r.get('subdir') or '')
        if not all(re.match(r'^[^\\/:*?"<>|]+$', seg) for seg in subdir.split('/') if seg):
            return jsonify(error='subdir 含非法字符'), 400
        xiaoban = (r.get('xiaoban') or '').strip()[:64]
        # 手机端 MediaRecorder 录制压缩后上传（precompressed=1）→ 服务器跳过转码；
        # 后缀与真实容器一致（Chrome Android 多为 mp4/h264，旧 WebView 为 webm）
        ext = (r.get('ext') or '').lower().strip('.')
        if ext not in ('mp4', 'webm', 'mov'):
            ext = 'mp4'
        precompressed = (r.get('precompressed') or '') == '1'
        fname = _video_name(r.get('filename') or '') + '.' + ext
        qiniu_key, remote_path = sync_cloud.photo_paths(s, subdir, fname)
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        cur = con.execute(
            'INSERT INTO photo_sync(workbook_id, xiaoban, filename, remote_path, qiniu_key,'
            ' size, state, created_at, kind, orig_size) VALUES(?,?,?,?,?,?,?,?,?,?)',
            (wid, xiaoban, fname, remote_path, qiniu_key, 0, 'received', now, 'video', len(data)))
        pid = cur.lastrowid
        con.commit()
        (PENDING_DIR / f'{pid}.mp4').write_bytes(data)
        threading.Thread(target=_process_video, args=(pid, precompressed), daemon=True).start()
        return jsonify(ok=True, id=pid, state='received',
                       note='视频已接收，云端压缩同步中' if not precompressed else '视频已接收，同步中')
    finally:
        con.close()


# ────────────────────────── 导出（按上传模板回填，保留原始格式） ──────────────────────────
# C07：相片不落服务器，原相片上传/预览/zip 路由已移除（v0.8）。

_XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
_XLSM_MIME = 'application/vnd.ms-excel.sheet.macroEnabled.12'


_DATE_ONLY_RE = re.compile(r'^(\d{4})-(\d{1,2})-(\d{1,2})$')
_DATETIME_RE = re.compile(r'^(\d{4})-(\d{1,2})-(\d{1,2})[ T](\d{1,2}):(\d{2})(?::(\d{2}))?$')


def _date_typed(text):
    """日期文本 → (datetime, Excel 数字格式)；不是日期则 (None, None)。

    只认 'YYYY-MM-DD' 与 'YYYY-MM-DD HH:MM[:SS]'（App 端与 Excel 日期单元格都长这样）。
    """
    s = ('' if text is None else str(text)).strip()
    m = _DATE_ONLY_RE.match(s)
    if m:
        try:
            return datetime(int(m[1]), int(m[2]), int(m[3])), 'yyyy-mm-dd'
        except ValueError:
            return None, None
    m = _DATETIME_RE.match(s)
    if m:
        try:
            return (datetime(int(m[1]), int(m[2]), int(m[3]),
                             int(m[4]), int(m[5]), int(m[6] or 0)), 'yyyy-mm-dd hh:mm')
        except ValueError:
            return None, None
    return None, None


def _export_workbook(wid, filters=None):
    """用上传时留存的原始模板（source.xlsx/xlsm）回填数据行，保留模板全部格式。
    v0.25：filters 为 {字段名: 值}，按「导出筛选」声明只筛行（不筛列）。
    返回 (BytesIO, 下载文件名, mimetype)；模板缺失 404。"""
    con = db()
    r = con.execute('SELECT * FROM workbooks WHERE id=?', (wid,)).fetchone()
    con.close()
    if not r:
        abort(404)
    ext = '.xlsm' if r['name'].lower().endswith('.xlsm') else '.xlsx'
    src = UPLOAD_DIR / str(wid) / f'source{ext}'
    if not src.exists():
        abort(404, description='原始模板文件缺失（该工作簿为旧版上传），请重新上传后导出')
    headers = json.loads(r['headers'])
    rows = json.loads(r['rows'])
    cfg = json.loads(r['config'])
    # v0.25：按参数「导出筛选」声明的字段筛行（只筛行、不筛列；未声明的字段一律忽略）
    rows = filter_rows(rows, headers, filters, dict(export_filters_of(headers, cfg)))
    wb = openpyxl.load_workbook(src, keep_vba=(ext == '.xlsm'))
    if r['sheet_name'] not in wb.sheetnames:
        abort(500, description=f'模板中找不到数据 sheet「{r["sheet_name"]}」')
    ws = wb[r['sheet_name']]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)   # 清掉模板里的示例/旧数据行，表头及格式保留
    ncols = len(headers)

    # 找出"整列都是日期文本"的列 → 导出成真日期。
    # 目的：Excel 自动筛选对**真日期列**才给「日期筛选」（年/月/日 树），
    # 纯文本日期只能给文本筛选，用户要的"日期选择"就出不来。
    date_cols = set()
    for i in range(ncols):
        vals = [str(row[i]) for row in rows if i < len(row) and str(row[i] or '').strip()]
        if vals and all(_date_typed(v)[0] for v in vals):
            date_cols.add(i)

    for row in rows:
        src_row = (list(row) + [''] * ncols)[:ncols]
        ws.append(src_row)
        n = ws.max_row
        for i in date_cols:
            dt, fmt = _date_typed(str(src_row[i]))
            if dt is not None:
                c = ws.cell(n, i + 1)
                c.value = dt
                c.number_format = fmt

    # 自动筛选范围按**实际数据**重设：
    # 模板自带的范围是上传时的旧值（行数一变就漏），源模板也可能压根没设。
    ws.auto_filter.ref = f'A1:{get_column_letter(ncols)}{ws.max_row}'
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stem = re.sub(r'\.xlsxm?$', '', r['name'], flags=re.IGNORECASE)
    return buf, f'{stem}_导出{ext}', _XLSM_MIME if ext == '.xlsm' else _XLSX_MIME


@bp.route('/api/workbooks/<int:wid>/export', methods=['GET'])
@login_required
def api_export(wid):
    """前端导出：登录用户即可按上传模板导出当前数据（已下架的工作簿 user 端导不出，C13）。"""
    con = db()
    inactive = _inactive_for(con, wid)
    con.close()
    if inactive:
        return jsonify(error=INACTIVE_MSG), 403
    buf, name, mime = _export_workbook(wid)
    return send_file(buf, as_attachment=True, download_name=name, mimetype=mime)


# ────────────────────────── 管理后台 API ──────────────────────────

@bp.route('/admin/api/workbooks', methods=['GET'])
@admin_required
def admin_list():
    con = db()
    rows = con.execute(
        'SELECT id, name, sheet_name, uploaded_at, config, key_column, headers,'
        ' COALESCE(is_active, 1) AS is_active FROM workbooks ORDER BY id DESC').fetchall()
    baks = {b['workbook_id']: b['n'] for b in con.execute(
        'SELECT workbook_id, COUNT(*) AS n FROM workbook_backups GROUP BY workbook_id')}
    con.close()
    out = []
    for r in rows:
        cfg = json.loads(r['config'])
        out.append({
            'id': r['id'], 'name': r['name'], 'sheet_name': r['sheet_name'],
            'uploaded_at': r['uploaded_at'],
            'is_active': bool(r['is_active']),      # C13：后台默认全列，前端自行按此过滤
            # 参数摘要：让管理端直接看到该 Excel 决定了哪些项可编辑 / 启用了什么功能
            # （config 中列表类 key 存的是分号分隔的原始字符串，这里归一化为数组）
            'editable_cols': split_list(cfg.get('可编辑列')),
            'hidden_cols': split_list(cfg.get('不显示列')),
            'features': split_list(cfg.get('功能')),
            'result_options': split_list(cfg.get('验收结果选项')),
            'unique_key': unique_key_of(cfg),      # v0.23：该工作簿声明的唯一键列（未声明则为空）
            'log_fields': split_list(cfg.get('日志字段')),
            # v0.25：后台「下载 Excel」可用的筛选字段（声明 ∩ 表头；空数组＝该模板没配筛选）
            'export_filters': [f for f, _k in export_filters_of(json.loads(r['headers']), cfg)],
            # v0.26：是否有「更新前」备份可回滚
            'can_rollback': baks.get(r['id'], 0) > 0,
        })
    return jsonify(workbooks=out)


@bp.route('/admin/api/workbooks/<int:wid>/download', methods=['GET'])
@admin_required
def admin_download(wid):
    """后台下载：与前端导出同链路（按上传模板回填，保留格式）。

    v0.25：支持 `?<字段>=<值>` 形式的筛选，只筛行。**只认参数「导出筛选」里声明过、
    且该工作簿表头确实存在的字段** —— 其它参数一律忽略（避免多传参数就把数据筛空）。
    """
    con = db()
    r = con.execute('SELECT headers, config FROM workbooks WHERE id=?', (wid,)).fetchone()
    con.close()
    if not r:
        abort(404)
    headers = json.loads(r['headers'])
    declared = [f for f, _k in export_filters_of(headers, json.loads(r['config']))]
    filters = {f: request.args.get(f) for f in declared if request.args.get(f)}
    buf, name, mime = _export_workbook(wid, filters)
    return send_file(buf, as_attachment=True, download_name=name, mimetype=mime)


@bp.route('/admin/api/workbooks/<int:wid>/filter-options', methods=['GET'])
@admin_required
def admin_filter_options(wid):
    """导出筛选弹框的字段与候选值。

    字段清单 = 参数「导出筛选」声明 ∩ 该工作簿表头（**没有的字段不返回**，弹框也就不显示）。
    select 类型附候选值（三路合并：参数枚举 / 系统用户 / 数据里出现过的值，见 param_parser.filter_options）；
    date 类型不需要候选值。
    """
    con = db()
    r = con.execute('SELECT headers, rows, config FROM workbooks WHERE id=?', (wid,)).fetchone()
    users = [x['username'] for x in
             con.execute('SELECT username FROM users ORDER BY role DESC, username')]
    con.close()
    if not r:
        abort(404)
    headers = json.loads(r['headers'])
    rows = json.loads(r['rows'])
    cfg = json.loads(r['config'])
    fields = []
    for field, kind in export_filters_of(headers, cfg):
        fields.append({
            'field': field,
            'type': kind,
            'options': filter_options(rows, headers, field, cfg, users) if kind == 'select' else [],
        })
    return jsonify(fields=fields)


@bp.route('/admin/api/accept-logs', methods=['GET'])
@admin_required
def admin_accept_logs():
    """验收字段变更日志（最近 300 条，新在前）。"""
    con = db()
    rows = con.execute(
        'SELECT l.id, l.workbook_id, w.name AS wb_name, l.row_idx, l.xiaoban,'
        ' l.field, l.old_value, l.new_value, l.operator, l.created_at'
        ' FROM accept_logs l LEFT JOIN workbooks w ON w.id = l.workbook_id'
        ' ORDER BY l.id DESC LIMIT 300').fetchall()
    con.close()
    return jsonify(logs=[dict(x) for x in rows])


# ────────────────────────── APK 直链（手机扫码/直接下载安装） ──────────────────────────

APK_DIR = BASE / 'static'   # 由 deploy.sh 同步，CI 产物放这里


@bp.route('/apk')
def apk_latest():
    """最新 Android 调试包直链（短链，便于二维码/手机输入）：
    取 static/supervision-*.apk 中文件名最大者（带版本号，天然按版本排序）。"""
    files = sorted(APK_DIR.glob('supervision-*.apk'))
    if not files:
        abort(404, description='尚无 APK 产物（CI 打包后放入 static/）')
    f = files[-1]
    return send_file(f, as_attachment=True, download_name=f.name,
                     mimetype='application/vnd.android.package-archive')


@bp.route('/guide')
def guide():
    """使用说明在线版（手机可直接打开/转发；不含账号密码表）。"""
    f = BASE / 'static' / 'guide.html'
    if not f.exists():
        abort(404)
    return send_file(f)


# ────────────────────────── 后台：同步设置与状态（doc/007 §5/§7） ──────────────────────────

_SYNC_EDITABLE = (set(sync_cloud.SYNC_DEFAULTS) | set(sync_cloud.VIDEO_DEFAULTS)
                  | set(sync_cloud.PHOTO_DEFAULTS))  # 允许后台写入的键（压缩与同步都可改）
_SYNC_SECRET = sync_cloud.SECRET_KEYS


def _mask(s: dict) -> dict:
    out = dict(s)
    out['has_qiniu_sk'] = bool(s.get('sync_qiniu_sk'))
    out['has_baidu_secret'] = bool(s.get('sync_baidu_secret_key'))
    tok = sync_cloud._load_token(s)
    out['has_token'] = bool(tok.get('access_token'))
    out['token_expires_at'] = (int(tok['expires_at']) if tok.get('expires_at') else None)
    for k in _SYNC_SECRET:
        out.pop(k, None)
    return out


@bp.route('/admin/api/sync/settings', methods=['GET'])
@admin_required
def admin_sync_get():
    con = db()
    s = sync_cloud.get_settings(con)
    con.close()
    ready, missing = sync_cloud.sync_ready(s)
    return jsonify(settings=_mask(s), ready=ready, missing=missing)


@bp.route('/admin/api/sync/settings', methods=['POST'])
@admin_required
def admin_sync_set():
    data = request.get_json(force=True)
    con = db()
    try:
        for k, v in (data.get('settings') or {}).items():
            if k not in _SYNC_EDITABLE:
                return jsonify(error=f'未知配置键：{k}'), 400
            v = str(v).strip()
            if k in _SYNC_SECRET and not v:
                continue  # secret 留空 = 保持原值
            if k == 'sync_baidu_token' and v:
                try:
                    tok = json.loads(v)
                    assert tok.get('access_token') and tok.get('refresh_token')
                except Exception:
                    return jsonify(error='token JSON 无效：需含 access_token 与 refresh_token'), 400
                tok.setdefault('expires_at', __import__('time').time() + tok.get('expires_in', 2592000))
                v = json.dumps(tok, ensure_ascii=False)
            if k == 'sync_enabled' and v == '1':
                s = sync_cloud.get_settings(con)
                probe = dict(s)
                probe.update({kk: str(vv).strip() for kk, vv in (data.get('settings') or {}).items()})
                ok, missing = sync_cloud.sync_ready(probe)
                if not ok:
                    return jsonify(error='开启失败，缺配置：' + '、'.join(missing)), 400
            sync_cloud.set_setting(con, k, v)
        con.commit()
        s = sync_cloud.get_settings(con)
        return jsonify(ok=True, settings=_mask(s))
    finally:
        con.close()


@bp.route('/admin/api/sync/status', methods=['GET'])
@admin_required
def admin_sync_status():
    con = db()
    try:
        counts = {r['state']: r['n'] for r in con.execute(
            'SELECT state, COUNT(*) AS n FROM photo_sync GROUP BY state')}
        recent = con.execute(
            'SELECT p.id, p.workbook_id, w.name AS wb_name, p.xiaoban, p.filename,'
            ' p.state, p.retry_count, p.last_error, p.created_at, p.synced_at,'
            ' p.kind, p.size, p.orig_size'
            ' FROM photo_sync p LEFT JOIN workbooks w ON w.id = p.workbook_id'
            ' ORDER BY p.id DESC LIMIT 30').fetchall()
        purged = sync_cloud.purge_expired(con)  # 顺手做生命周期清理
        s = sync_cloud.get_settings(con)
        return jsonify(counts=counts, purged=purged,
                       recent=[dict(x) for x in recent],
                       token_expires_at=_mask(s).get('token_expires_at'))
    finally:
        con.close()


@bp.route('/admin/api/sync/retry', methods=['POST'])
@admin_required
def admin_sync_retry():
    ids = (request.get_json(silent=True) or {}).get('ids') or []
    con = db()
    try:
        if ids:
            rows = con.execute(
                "SELECT * FROM photo_sync WHERE state IN ('received','qiniu_ok')"
                " AND id IN (%s)" % ','.join('?' * len(ids)), ids).fetchall()
        else:
            rows = con.execute(
                "SELECT * FROM photo_sync WHERE state IN ('received','qiniu_ok')"
                " ORDER BY id LIMIT 100").fetchall()
        ok = fail = 0
        for row in rows:
            try:
                sync_cloud.push_to_baidu(con, row, PENDING_DIR)
                con.commit()
                ok += 1
            except sync_cloud.SyncError as e:
                con.execute("UPDATE photo_sync SET retry_count=retry_count+1, last_error=?"
                            " WHERE id=?", (str(e)[:500], row['id']))
                con.commit()
                fail += 1
        return jsonify(ok=ok, fail=fail)
    finally:
        con.close()


@bp.route('/admin/api/tracks', methods=['GET'])
@admin_required
def admin_tracks():
    files = sorted(TRACK_DIR.glob('*.gpx'), reverse=True)
    return jsonify(tracks=[{'file': f.name, 'size': f.stat().st_size} for f in files])


@bp.route('/admin/api/tracks.zip', methods=['GET'])
@admin_required
def admin_tracks_zip():
    """轨迹压缩包（v0.20）：**每小班一个文件夹**，内含 shapefile 组件（.shp/.shx/.dbf/.prj/.cpg）
    与原始 GPX；WGS84 + UTF-8，ArcGIS 10.1+/QGIS 直接打开（口径对齐 hqz-survey R17）。"""
    try:
        buf, stats = track_export.export_tracks_zip(TRACK_DIR)
    except ValueError as e:
        return jsonify(error=str(e)), 404
    name = '轨迹导出_shp_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '.zip'
    resp = send_file(buf, as_attachment=True, download_name=name, mimetype='application/zip')
    resp.headers['X-Track-Export'] = json.dumps(stats, ensure_ascii=False)
    return resp


@bp.route('/admin/api/tracks/<path:name>', methods=['GET'])
@admin_required
def admin_track_dl(name):
    path = (TRACK_DIR / Path(name).name)
    if not path.exists():
        abort(404)
    return send_file(path, as_attachment=True, download_name=path.name)


def create_app(prefix=''):
    """Flask 工厂。gateway 约定 create_app(prefix=spec['prefix'])；
    本地预览 prefix=''（路由在根路径）。
    注意：gateway DispatcherMiddleware 已设 SCRIPT_NAME 剥前缀，路由不挂 url_prefix
    （否则双重前缀 404）；static 用默认 /static 规则（剥离后正好命中）；
    prefix 仅用于 session cookie 限定。
    """
    app = Flask(__name__)
    app.secret_key = 'hqz-supervision-local-preview-key'
    # 会话长期有效（v0.13.2）：系统相机/相册等外部 Activity 会短暂顶掉 WebView，
    # 若 cookie 是"会话级"，WebView 重启即丢 → 页面重载后要求重新登录（2026-09-13 实机反馈）。
    # 改为 30 天持久 cookie；生产（有 prefix）额外开启 Secure。
    app.permanent_session_lifetime = timedelta(days=30)
    if prefix:
        app.config['SESSION_COOKIE_SECURE'] = True
        app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'

    @app.after_request
    def _no_cache_html(resp):
        # HTML 不缓存：发版后浏览器必须拿新页面（静态资源 /static 不受限，正常缓存）
        if resp.mimetype == 'text/html':
            resp.headers['Cache-Control'] = 'no-cache'
        return resp

    if prefix:
        app.config['APPLICATION_ROOT'] = prefix      # session cookie 限定本应用前缀
        app.config['PREFERRED_URL_SCHEME'] = 'https'
    app.register_blueprint(bp)
    return _EmptyPathFix(app)


class _EmptyPathFix:
    """网关 DispatcherMiddleware 剥前缀后 PATH_INFO 可能为空串（如 GET /supervision），
    werkzeug 会 308「补斜杠」且 Location 用 environ 的明文 scheme 拼成
    http://forest.bibook.top/supervision/ —— Android WebView 直接
    net::ERR_CLEARTEXT_NOT_PERMITTED（2026-09-11 v0.7.2 APK 实测踩坑）。
    空路径归一为 /，直接命中首页，绕开重定向（浏览器也少一跳）。
    nginx 在容器内无权改，故在应用层根治；hqz-survey 同架构同坑（其用
    usesCleartextTraffic=true 掩盖，本项目不走明文）。"""

    def __init__(self, app):
        self.app = app

    def __getattr__(self, name):
        # gateway mount 会对工厂返回值访问 .config 等属性，必须透传给 Flask app
        return getattr(self.app, name)

    def __call__(self, environ, start_response):
        if environ.get('PATH_INFO', '') == '':
            environ['PATH_INFO'] = '/'
        return self.app(environ, start_response)


init_db()

if __name__ == '__main__':
    print(f'hqz-supervision 本地预览 → http://127.0.0.1:{PORT} （不影响其他业务）')
    create_app().run(host='127.0.0.1', port=PORT, debug=False)
