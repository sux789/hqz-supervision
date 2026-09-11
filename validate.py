"""输入校验骨架：坏数据入口炸掉，报错带行号+字段名+当前值+可选值。
按项目改下方三张表（REQUIRED / UNIQUE_KEYS / ALLOWED）即可用。
用法: python validate.py data/xxx.xlsx ；全部通过退出码 0，否则 1。
WARN 不阻断，汇总输出；FATAL 阻断。所有记录被跳过 = FATAL。"""
import sys
from pathlib import Path

# ── 按项目填写 ──────────────────────────────────────
SHEET = None                                    # 校验哪个 sheet，None=第一个
HEADER_ROW = 1                                  # 表头所在行
REQUIRED = ["关键列1", "关键列2"]                # 必填表头（标准名）
ALIASES = {"关键列1": ["关键列1", "别名1"],
           "关键列2": ["关键列2", "别名2"]}      # 标准名 → 可接受的表头列表
UNIQUE_KEYS = ["关键列1"]                        # 联合唯一键，重复即 FATAL
ALLOWED = {"关键列2": ["可选值A", "可选值B"]}     # 参考表：值不在表内原样报错+列可选值
# ────────────────────────────────────────────────────

errors, warns = [], []

def fail(row, field, value, msg):
    errors.append(f"行{row} | {field}={value!r} | {msg}")

def validate_rows(header: list, rows: list):
    # 表头别名归一：标准名 → 实际列号
    col = {}
    for std, names in ALIASES.items():
        for n in names:
            if n in header:
                col[std] = header.index(n)
                break
    for req in REQUIRED:
        if req not in col:
            errors.append(f"表头缺关键列 {req!r}（可接受别名: {ALIASES.get(req)}）")
    if errors:
        return
    seen, skipped = {}, 0
    for r, row in enumerate(rows, start=HEADER_ROW + 1):
        if all(v is None or str(v).strip() == "" for v in row):
            skipped += 1
            continue  # 整行空白跳过（计数，最后判定）
        for req in REQUIRED:
            v = row[col[req]]
            if v is None or str(v).strip() == "":
                fail(r, req, v, "必填为空")
        key = tuple(row[col[k]] for k in UNIQUE_KEYS if k in col)
        if key in seen:
            fail(r, "+".join(UNIQUE_KEYS), key, f"与行{seen[key]} 重复，唯一键冲突")
        else:
            seen[key] = r
        for f, allowed in ALLOWED.items():
            if f in col:
                v = row[col[f]]
                if v is not None and str(v).strip() != "" and v not in allowed:
                    fail(r, f, v, f"不在参考表，可选值: {allowed}")
    total = len(rows)
    if total > 0 and skipped == total:
        errors.append(f"所有 {total} 行记录均被跳过（整行空白），禁止静默生成空输出")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    p = Path(sys.argv[1])
    if p.suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(p, data_only=True)
        ws = wb[SHEET] if SHEET else wb.worksheets[0]
        data = [[c.value for c in row] for row in ws.iter_rows()]
    else:
        data = [line.split(",") for line in p.read_text(encoding="utf-8").splitlines()]
    if not data:
        errors.append("输入为空")
    else:
        validate_rows(data[HEADER_ROW - 1], data[HEADER_ROW:])
    for w in warns:
        print(f"[WARN] {w}")
    for e in errors:
        print(f"[FATAL] {e}")
    print(f"\n校验: {len(errors)} 错误 / {len(warns)} 警告")
    sys.exit(1 if errors else 0)
